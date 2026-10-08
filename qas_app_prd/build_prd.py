from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from pathlib import Path

OUT_DIR = Path(r"C:\Users\Ari\.codex\.chatgpt-projects\g-p-6aacf1deeee08191b49d027b1363f153\outputs\qas_app_prd")
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "PRD_Aplikasi_Audit_QAS_Cloudflare.docx"

NAVY = "0B2D57"
NAVY_2 = "173B67"
GREEN = "70AD47"
LIGHT_GREEN = "E2F0D9"
CREAM = "F7EBD2"
LIGHT_BLUE = "EAF1F8"
LIGHT_GREY = "F2F4F7"
MID_GREY = "D9DEE7"
DARK = "17243A"
WHITE = "FFFFFF"
AMBER = "F4B183"
RED = "C00000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    total_twips = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    for row in table.rows:
        for i, width in enumerate(widths):
            cell = row.cells[i]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def keep_table_together(table):
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                set_keep_with_next(paragraph, True)
    if table.rows:
        last_cell = table.rows[-1].cells[-1]
        if last_cell.paragraphs:
            p_pr = last_cell.paragraphs[-1]._p.get_or_add_pPr()
            keep = p_pr.find(qn("w:keepNext"))
            if keep is not None:
                p_pr.remove(keep)


def set_cell_text(cell, text, bold=False, color=DARK, size=8.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    return p


def add_hyperlink(paragraph, text, url, color=NAVY, underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(fonts)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_keep_with_next(paragraph, value=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if value and keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_prefix=None, italic=False):
    p = doc.add_paragraph(style="Body Text")
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    if italic:
        for r in p.runs:
            r.italic = True
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def add_callout(doc, title, body, fill=LIGHT_BLUE, accent=NAVY):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_widths(table, [0.15, 7.0])
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    p = table.cell(0, 1).paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(accent)
    r2 = p.add_run(body)
    r2.font.name = "Arial"
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor.from_string(DARK)
    set_cell_margins(table.cell(0, 1), 120, 150, 120, 150)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths=None, header_fill=NAVY, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_shading(hdr.cells[i], header_fill)
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        if r_idx % 2:
            for c in cells:
                set_cell_shading(c, LIGHT_GREY)
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    body = styles["Body Text"]
    body.font.name = "Arial"
    body._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    body._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    body.font.size = Pt(9.5)
    body.font.color.rgb = RGBColor.from_string(DARK)
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_after = Pt(14)

    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 12.5, GREEN), ("Heading 3", 10.5, NAVY_2)):
        s = styles[name]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(5)

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        s = styles[name]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        s.font.size = Pt(9.3)
        s.paragraph_format.space_after = Pt(3)

    if "Code" not in styles:
        code = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code"]
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor.from_string(DARK)
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.right_indent = Inches(0.25)
    code.paragraph_format.space_after = Pt(6)


def configure_section(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    hp = section.header.paragraphs[0]
    hp.text = "PRD Aplikasi Audit QAS"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    fp = section.footer.paragraphs[0]
    fp.text = "Dokumen produk dan implementasi | 1 Oktober 2026   "
    for r in fp.runs:
        r.font.name = "Arial"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("667085")
    add_page_number(fp)


questions = [
    {
        "id": "J1-DIS-01", "group": "J1 AHM to Main Dealer", "section": "Distribusi unit SMH dan asesoris dari AHM to Main Dealer", "score": "4", "mandatory": "Ya",
        "question": "Apakah telah dilakukan kontrol oleh Main Dealer untuk proses transportasi dari AHM ke Main Dealer? (MD)",
        "options": [
            "Dilakukan review ekspedisi rutin tiap periodik",
            "Ekspedisi hanya direview setiap kali ada masalah, seperti NRFS tinggi, keterlambatan, dan sebagainya",
            "Ekspedisi dikontrol oleh Head Office; Main Dealer tidak mengetahui proses kontrol",
            "Tidak ada kontrol dari Main Dealer",
        ],
    },
    {
        "id": "J1-DIS-02", "group": "J1 AHM to Main Dealer", "section": "Distribusi unit SMH dan asesoris dari AHM to Main Dealer", "score": "4", "mandatory": "Ya",
        "question": "Bagaimana proses unloading dilakukan, dan apakah unit yang masuk ke gudang disertai dokumen terkait? (MD)",
        "options": [
            "Proses unloading dilakukan sesuai prosedur lengkap dengan shipping list dan list accessory",
            "Unit diturunkan sesuai prosedur tetapi dokumen shipping list dan accessory tidak lengkap",
            "Unit diturunkan dengan proses yang berpotensi membahayakan unit dan dokumen lengkap",
            "Ada improvement terkait proses unloading atau bongkar muat",
        ],
    },
    {
        "id": "J1-DIS-03", "group": "J1 AHM to Main Dealer", "section": "Distribusi unit SMH dan asesoris dari AHM to Main Dealer", "score": "4", "mandatory": "Ya",
        "question": "Apakah dilakukan pemeriksaan Unit Motor secara visual dan kelengkapan part saat akan masuk ke gudang?",
        "options": [
            "Dilakukan pengecekan unit secara visual saja",
            "Dilakukan pengecekan unit secara visual dan kelengkapan ACC",
            "Tidak dilakukan pengecekan",
            "Dilakukan pengecekan visual dan kelengkapan ACC serta didokumentasikan",
        ],
    },
    {
        "id": "J1-NRFS-01", "group": "J1 AHM to Main Dealer", "section": "Penanganan unit SMH Not Ready For Sale", "score": "5", "mandatory": "Ya",
        "question": "Apa yang dilakukan jika menemukan unit cacat atau lecet akibat transportasi?",
        "options": [
            "Lapor ke sopir atau deliveryman dan admin; unit ditempatkan di area khusus NRFS, diberi identifikasi, dicatat kerusakannya, dan dilaporkan ke Main Dealer",
            "Lapor ke sopir atau deliveryman dan admin; unit cacat atau lecet disimpan di gudang sesuai tipe dan warna",
            "Lapor ke sopir atau deliveryman dan admin; lanjut memeriksa unit lain tanpa tindakan pada unit cacat atau lecet",
            "Lapor ke sopir atau deliveryman dan admin; unit dikembalikan ke gudang Main Dealer untuk ditukar",
            "Ada improvement terkait penanganan unit cacat atau lecet akibat transportasi",
        ],
    },
    {
        "id": "J1-NRFS-02", "group": "J1 AHM to Main Dealer", "section": "Penanganan unit SMH Not Ready For Sale", "score": "4", "mandatory": "Ya",
        "question": "Apakah tersedia lokasi khusus untuk unit lecet, cacat, atau not good yang Not Ready for Sale?",
        "options": [
            "Ada lokasi NRFS tetapi tidak ada tanda atau identifikasi",
            "Tidak ada lokasi khusus NRFS atau bergabung dengan unit lain",
            "Ada lokasi khusus NRFS yang memiliki tanda atau identifikasi",
        ],
    },
    {
        "id": "J1-NRFS-03", "group": "J1 AHM to Main Dealer", "section": "Penanganan unit SMH Not Ready For Sale", "score": "4", "mandatory": "Ya",
        "question": "Apakah terdapat PIC yang bertugas memperbaiki unit NRFS?",
        "options": [
            "Tidak ada PIC yang melakukan repair",
            "Ada PIC repair yang sudah training minimal TTL 1",
            "Ada PIC repair yang sudah training minimal TTL 2",
            "Ada PIC repair tetapi belum training",
            "Ada PIC repair yang sudah training TTL 3",
        ],
    },
    {
        "id": "J1-NRFS-04", "group": "J1 AHM to Main Dealer", "section": "Penanganan unit SMH Not Ready For Sale", "score": "4", "mandatory": "Tidak ditandai pada PDF",
        "question": "Bagaimana proses penggantian part unit NRFS jika tidak memiliki stok part?",
        "options": [
            "Beli di H3 terdekat, part shop satelit, HEPS, atau AHASS terdekat",
            "Membeli part di bengkel umum atau toko spare part umum",
            "Urgent order ke AHM",
        ],
    },
    {
        "id": "J1-NRFS-05", "group": "J1 AHM to Main Dealer", "section": "Penanganan unit SMH Not Ready For Sale", "score": "4", "mandatory": "Ya",
        "question": "Siapa yang melakukan pengecekan dan memastikan unit NRFS selesai direpair dan dinyatakan OK?",
        "options": [
            "Mekanik yang memperbaiki unit itu sendiri dan sudah lulus minimal TTL 1",
            "Mekanik lain dengan kompetensi setara atau lebih tinggi dari mekanik yang melakukan perbaikan",
            "Tidak ada; setelah diperbaiki unit otomatis masuk stok unit OK",
            "Kepala gudang yang sedang dalam pengajuan training TTL",
        ],
    },
    {
        "id": "J1-NRFS-06", "group": "J1 AHM to Main Dealer", "section": "Penanganan unit SMH Not Ready For Sale", "score": "5", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apa yang dilakukan oleh Main Dealer untuk memonitor penyelesaian unit NRFS?",
        "options": [
            "Monitoring order part ke AHM secara berkala dan terdata",
            "Tidak melakukan tindakan",
            "Menanyakan status kedatangan part ke AHM secara berkala tetapi tidak terdata",
            "Ada improvement terkait monitoring penyelesaian unit NRFS",
        ],
    },
    {
        "id": "J1-MNT-01", "group": "J1 AHM to Main Dealer", "section": "Maintenance Unit SMH di Gudang Main Dealer", "score": "5", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apakah ada unit yang disimpan lebih dari satu bulan dan dilakukan pengecekan serta perawatan, termasuk tekanan ban, debu, karat, rantai, dan jarak antarunit?",
        "options": [
            "Ada; unit dibersihkan setiap hari tetapi tidak didokumentasikan",
            "Tidak ada unit SMH di gudang yang berumur lebih dari satu bulan",
            "Ada tetapi tidak dilakukan pengecekan dan perawatan",
            "Ada; pengecekan dan perawatan dilakukan tetapi tidak konsisten dan tidak terjadwal",
            "Ada; pengecekan dan perawatan dilakukan rutin, terjadwal, dan terdokumentasi",
            "Ada improvement terhadap aktivitas perawatan unit SMH di gudang",
        ],
    },
    {
        "id": "J1-MNT-02", "group": "J1 AHM to Main Dealer", "section": "Maintenance Unit SMH di Gudang Main Dealer", "score": "5", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apa yang dilakukan jika menemukan unit NG, cacat, atau lecet ketika disimpan, dikeluarkan, atau dipindahkan di area gudang?",
        "options": [
            "Tidak pernah ditemukan unit NG di gudang unit dealer atau display",
            "Langsung diperbaiki sendiri oleh orang yang menemukan",
            "Langsung dilaporkan kepada pihak berwenang; unit tetap di lokasi semula",
            "Dipindahkan ke lokasi khusus, diberi identifikasi NRFS, dibuatkan laporan, dan dilaporkan kepada pihak berwenang",
            "Dipindahkan ke lokasi NRFS dan diberi identifikasi bagian yang cacat",
            "Ada improvement terhadap aktivitas penanganan unit cacat akibat penyimpanan",
        ],
    },
    {
        "id": "J1-MNT-03", "group": "J1 AHM to Main Dealer", "section": "Maintenance Unit SMH di Gudang Main Dealer", "score": "4", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apakah ada tindakan yang dilakukan terhadap penyebab unit NG akibat penyimpanan?",
        "options": [
            "Selalu ada tindakan perbaikan terhadap penyebab unit NG akibat simpan",
            "Ada rencana perbaikan tetapi belum dilaksanakan karena menunggu persetujuan",
            "Tidak ada tindakan karena penyebab unit NG tidak diketahui",
            "Tidak pernah ada kasus unit NG akibat simpan",
        ],
    },
    {
        "id": "J2-DIS-01", "group": "J2 Main Dealer to Dealer", "section": "Distribusi unit SMH dan asesoris dari Main Dealer to Dealer", "score": "5", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apa yang dilakukan oleh dealer untuk memastikan battery yang diterima dari AHM selalu dalam kondisi bagus?",
        "options": [
            "Mengirimkan memo, WhatsApp, atau pemberitahuan lisan kepada AHM agar mengirim battery terbaru",
            "Melakukan pemeriksaan umur semua battery pada setiap penerimaan dari AHM",
            "Tidak melakukan tindakan",
            "Ada improvement terhadap proses untuk memastikan battery yang diterima selalu dalam kondisi bagus",
        ],
    },
    {
        "id": "J2-DIS-02", "group": "J2 Main Dealer to Dealer", "section": "Distribusi unit SMH dan asesoris dari Main Dealer to Dealer", "score": "5", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apa yang dilakukan untuk memastikan Unit SMH yang dikirim ke dealer selalu FIFO?",
        "options": [
            "Pengambilan berdasarkan tag FIFO bulanan",
            "Pengambilan sesuai database unit ready for sale",
            "Pengambilan berdasarkan posisi paling depan, belakang, kanan, atau kiri",
            "Ada improvement yang dilakukan terkait implementasi FIFO",
        ],
    },
    {
        "id": "J2-DIS-03", "group": "J2 Main Dealer to Dealer", "section": "Distribusi unit SMH dan asesoris dari Main Dealer to Dealer", "score": "4", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apakah ekspedisi atau kendaraan dealer memiliki perlengkapan pengikatan unit SMH sesuai Prosedur atau Instruksi Kerja?",
        "options": [
            "Bentuk, ukuran, dan jumlah peralatan pengikatan sesuai prosedur untuk setiap unit SMH",
            "Bentuk, ukuran, dan jumlah peralatan pengikatan tidak sesuai prosedur untuk setiap unit SMH",
            "Tidak memiliki perlengkapan pengikatan unit SMH",
            "Hanya sebagian perlengkapan pengikatan yang sesuai prosedur",
            "Ada improvement penggunaan peralatan pengikatan yang lebih baik dari prosedur",
        ],
    },
    {
        "id": "J2-DIS-04", "group": "J2 Main Dealer to Dealer", "section": "Distribusi unit SMH dan asesoris dari Main Dealer to Dealer", "score": "4", "mandatory": "Tidak ditandai pada PDF",
        "question": "Apakah unit SMH dari Main Dealer ke Dealer telah terikat sesuai Quality Point dalam Prosedur atau Instruksi Kerja?",
        "options": [
            "Semua Quality Point telah sesuai prosedur atau instruksi kerja",
            "Hanya sebagian Quality Point yang sesuai",
            "Pengikatan dilakukan tetapi tidak sesuai Quality Point",
            "Ada improvement terhadap Quality Point pengikatan",
            "Sebagian titik pengikatan tidak sesuai prosedur tetapi unit dijamin tidak mengalami penurunan kualitas saat diantar",
        ],
    },
    {
        "id": "J2-DIS-05", "group": "J2 Main Dealer to Dealer", "section": "Distribusi unit SMH dan asesoris dari Main Dealer to Dealer", "score": "4", "mandatory": "Tidak ditandai pada PDF",
        "question": "Bagaimana Main Dealer memastikan unit SMH yang dikirim ke dealer selalu dalam kondisi baik?",
        "options": [
            "Menyediakan PIC kompeten untuk pengecekan fisik dan dokumen secara sampling",
            "Menyediakan PIC kompeten yang bertugas khusus melakukan pengecekan fisik dan dokumen 100 persen",
            "Menyediakan PIC kompeten untuk melakukan pengecekan fisik dan dokumen 100 persen",
            "Setelah loading, unit langsung dikirim tanpa pemeriksaan",
        ],
    },
]


doc = Document()
configure_styles(doc)
configure_section(doc.sections[0])

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("PRD Aplikasi Audit QAS")
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(18)
r = sub.add_run("Quality Assurance System untuk Gudang Karawang Baros dan Cirebon")
r.font.name = "Arial"
r.font.size = Pt(15)
r.font.color.rgb = RGBColor.from_string(GREEN)
r.bold = True

meta = doc.add_table(rows=5, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.style = "Table Grid"
meta_data = [
    ("Versi", "1.1"),
    ("Tanggal", "1 Oktober 2026"),
    ("Platform target", "Cloudflare dan Progressive Web App"),
    ("Pengguna", "Auditor QAS dan PIC QAS Gudang"),
    ("Sumber form", "Form Audit QAS.pdf, dua halaman"),
]
for i, (a, b) in enumerate(meta_data):
    set_cell_shading(meta.cell(i, 0), NAVY if i == 0 else LIGHT_BLUE)
    set_cell_text(meta.cell(i, 0), a, bold=True, color=WHITE if i == 0 else NAVY, size=9)
    set_cell_text(meta.cell(i, 1), b, size=9)
set_table_widths(meta, [1.45, 5.7])

doc.add_paragraph()
add_callout(
    doc,
    "Rekomendasi inti",
    "Bangun satu aplikasi responsive PWA yang berjalan pada ponsel dan desktop. Gunakan Cloudflare Worker dengan Static Assets untuk antarmuka dan API, D1 untuk data terstruktur, R2 untuk bukti foto, serta Cloudflare Access untuk autentikasi. Self audit diselesaikan dan dikunci sebelum audit resmi; hasil keduanya dibandingkan otomatis per pertanyaan, bagian, depo, dan periode.",
    fill=LIGHT_GREEN,
    accent=GREEN,
)
add_callout(
    doc,
    "Pembaruan versi 1.1",
    "Menambahkan spesifikasi lengkap Master Template, Master Bagian, Master Soal, daftar pilihan, nilai setiap pilihan, aturan tambah edit hapus, clone version, validasi, preview, simulasi skor, publish, snapshot historis, API, audit event, dan acceptance criteria.",
    fill=CREAM,
    accent=NAVY,
)

doc.add_paragraph()
p = doc.add_paragraph()
p.add_run("Status dokumen: ").bold = True
p.add_run("Siap digunakan sebagai acuan build di Google Antigravity setelah keputusan penilaian pada Bagian 15 dikonfirmasi.")

doc.add_page_break()

add_heading(doc, "Daftar Isi", 1)
toc_items = [
    "1 Ringkasan Eksekutif",
    "2 Step 1 Pemahaman Tujuan Aplikasi",
    "3 Step 2 Rekomendasi dari Awal sampai Selesai",
    "4 Step 3 Product Requirements Document",
    "5 Pengguna Peran dan Hak Akses",
    "6 Alur Proses Audit",
    "7 Kebutuhan Fungsional",
    "8 Struktur Halaman dan Pengalaman Pengguna",
    "9 Dashboard dan Perbandingan",
    "10 Aturan Penilaian",
    "11 Arsitektur dan Penyimpanan Data",
    "12 Model Data",
    "13 Kontrak API",
    "14 Keamanan Privasi dan Audit Trail",
    "15 Keputusan yang Wajib Dikonfirmasi",
    "16 Acceptance Criteria",
    "17 Rencana Implementasi Antigravity",
    "18 Pengujian Peluncuran dan Operasional",
    "19 Risiko dan Mitigasi",
    "20 Sumber",
    "Lampiran A Pemetaan Form Audit QAS",
]
add_bullets(doc, toc_items)

doc.add_page_break()

add_heading(doc, "1 Ringkasan Eksekutif", 1)
add_body(doc, "Aplikasi Audit QAS akan mengganti form audit manual menjadi proses digital yang dapat digunakan dari ponsel saat berada di area gudang dan dari desktop untuk review, analisis, serta administrasi. Sistem melayani tiga depo, yaitu Karawang, Baros, dan Cirebon. PIC QAS Gudang melakukan self audit. Auditor QAS melakukan audit resmi. Kedua hasil disimpan secara terpisah dan dibandingkan setelah audit resmi diserahkan.")
add_body(doc, "Form sumber berisi 17 pertanyaan dalam kelompok J1 AHM to Main Dealer dan J2 Main Dealer to Dealer. Setiap pertanyaan memiliki beberapa pilihan jawaban, nilai, serta area dokumentasi. Nilai akhir contoh pada PDF adalah 4,37 dengan kategori BAIK SEKALI. PDF tidak menampilkan pemetaan skor setiap pilihan, rumus agregasi, bobot, atau rentang kategori secara lengkap. Karena itu, mesin penilaian harus berbasis konfigurasi versi template.")
add_callout(doc, "Keputusan produk", "Gunakan mode audit buta sebagai default. Auditor dapat melihat bahwa self audit sudah selesai, tetapi tidak dapat melihat jawaban dan skor self audit sebelum audit resmi diserahkan. Cara ini mengurangi bias dan membuat perbandingan lebih dapat dipercaya.", fill=CREAM, accent=NAVY)

add_heading(doc, "Tujuan keberhasilan", 2)
add_bullets(doc, [
    "PIC dapat menyelesaikan self audit dari ponsel, menyimpan draft, mengunggah bukti, dan menyerahkan hasil.",
    "Auditor dapat menjalankan audit resmi dengan template yang sama tanpa dipengaruhi jawaban self audit.",
    "Sistem menghitung skor sesuai versi aturan yang berlaku dan menampilkan perbedaan self audit terhadap audit resmi.",
    "Manajemen dapat melihat status, skor, gap, tren, dan pertanyaan bermasalah untuk setiap depo.",
    "Semua perubahan penting memiliki jejak audit dan bukti tersimpan secara privat.",
])

add_heading(doc, "Batas MVP", 2)
add_body(doc, "MVP mencakup login, pengelolaan pengguna sederhana, template audit berversi, self audit, audit resmi, bukti foto atau dokumen, penguncian hasil, perbandingan, dashboard, ekspor, dan audit trail. Modul tindak lanjut temuan, notifikasi WhatsApp, tanda tangan elektronik lanjutan, dan integrasi sistem perusahaan ditempatkan pada fase berikutnya.")

add_heading(doc, "2 Step 1 Pemahaman Tujuan Aplikasi", 1)
add_body(doc, "Maksud aplikasi adalah membangun satu sistem QAS yang menyatukan persiapan audit, pelaksanaan audit, dokumentasi bukti, perhitungan nilai, dan analisis hasil lintas gudang. Aplikasi bukan hanya form digital. Nilai utamanya berasal dari alur dua tahap dan perbandingan yang konsisten.")
add_table(doc, ["Aspek", "Pemahaman kebutuhan"], [
    ["Lokasi", "Tiga depo: Karawang, Baros, dan Cirebon"],
    ["Peran", "PIC QAS Gudang dan Auditor QAS"],
    ["Self audit", "Dilakukan PIC sebelum audit resmi"],
    ["Audit resmi", "Dilakukan Auditor menggunakan template dan versi yang sama"],
    ["Perbandingan", "Self audit dibandingkan dengan audit resmi pada tingkat pertanyaan, bagian, total, depo, dan periode"],
    ["Perangkat", "Ponsel untuk input lapangan dan desktop untuk review serta dashboard"],
    ["Data", "Jawaban, skor, catatan dokumentasi, foto atau file bukti, status, pengguna, waktu, dan histori perubahan"],
    ["Output", "Dashboard, detail hasil, gap, tren, status audit, dan ekspor"],
], widths=[1.45, 5.7])

add_heading(doc, "Prinsip proses", 2)
add_numbered(doc, [
    "Auditor membuat atau mengaktifkan periode audit untuk depo tertentu.",
    "PIC mengisi self audit, melengkapi bukti, melakukan review, lalu menyerahkan hasil.",
    "Self audit dikunci agar tidak berubah setelah diserahkan, kecuali dibuka kembali dengan alasan dan jejak audit.",
    "Auditor menjalankan audit resmi. Detail self audit disembunyikan sampai audit resmi diserahkan.",
    "Sistem menghitung hasil dan membuka halaman perbandingan.",
    "Auditor memfinalkan audit. PIC dapat melihat dan mengakui hasil untuk depornya.",
])

add_heading(doc, "3 Step 2 Rekomendasi dari Awal sampai Selesai", 1)
add_heading(doc, "3.1 Pilihan platform", 2)
add_body(doc, "Bangun responsive Progressive Web App dengan React dan TypeScript. PWA dapat dibuka melalui browser, dipasang ke layar utama ponsel, memanfaatkan kamera untuk unggah bukti, dan tetap nyaman digunakan di desktop. Pendekatan ini menjaga biaya dan kompleksitas lebih rendah daripada membuat aplikasi Android dan iOS terpisah.")

add_heading(doc, "3.2 Stack yang direkomendasikan", 2)
add_table(doc, ["Lapisan", "Pilihan", "Alasan"], [
    ["Frontend", "React, TypeScript, Vite, PWA", "Satu basis kode untuk mobile dan desktop; mudah dibangun dan diuji di Antigravity"],
    ["UI", "Tailwind CSS dan komponen aksesibel", "Responsive, cepat dibangun, dan konsisten"],
    ["Backend", "Cloudflare Worker dengan Hono", "API ringan, berjalan di edge, dan dapat disatukan dengan static assets"],
    ["Database", "Cloudflare D1", "Relasional dan cocok untuk audit, jawaban, versi template, peran, serta laporan"],
    ["File bukti", "Cloudflare R2 private bucket", "Foto dan dokumen tidak membebani database; egress gratis"],
    ["Login", "Cloudflare Access dengan email OTP atau akun Cloudflare", "Tidak perlu menyimpan password aplikasi; cocok untuk tim internal kecil"],
    ["Otorisasi", "RBAC di D1", "Cloudflare Access membuktikan identitas; aplikasi menentukan peran dan depo"],
    ["Deployment", "Wrangler dan Git repository", "Deployment dapat diulang dan dilacak"],
], widths=[1.1, 2.1, 3.95])

add_heading(doc, "3.3 Penyimpanan data yang direkomendasikan", 2)
add_body(doc, "Gunakan D1 untuk data transaksi dan R2 untuk lampiran. Jangan menyimpan foto sebagai BLOB atau base64 di D1. D1 harus menyimpan metadata file, sedangkan file fisik berada di R2. Bucket R2 tetap privat; file hanya diakses melalui endpoint aplikasi setelah izin pengguna diperiksa.")
add_table(doc, ["Jenis data", "Lokasi", "Contoh"], [
    ["Data terstruktur", "D1", "Pengguna, depo, periode, template, pertanyaan, jawaban, skor, status, audit trail"],
    ["Lampiran", "R2", "Foto kondisi gudang, shipping list, daftar accessory, dokumen pendukung"],
    ["Draft offline", "IndexedDB pada perangkat", "Jawaban dan antrean unggah yang belum tersinkron"],
    ["Konfigurasi rahasia", "Cloudflare secrets", "AUD Access, kunci penandatanganan internal, konfigurasi lingkungan"],
], widths=[1.45, 1.4, 4.3])

add_heading(doc, "3.4 Kelayakan versi gratis", 2)
add_table(doc, ["Layanan", "Batas gratis saat dokumen dibuat", "Catatan desain"], [
    ["Workers", "100.000 request per hari; 10 ms CPU per invocation", "Cukup untuk penggunaan internal tiga depo jika query dan payload efisien"],
    ["D1", "5 juta row read per hari; 100.000 row write per hari; 5 GB total storage", "Gunakan indeks dan pagination. Free database memiliki batas 500 MB per database"],
    ["R2 Standard", "10 GB-month storage; 1 juta operasi Class A; 10 juta Class B per bulan", "Kompres foto dan batasi jumlah atau ukuran bukti"],
    ["Cloudflare Access", "$0 untuk tim di bawah 50 pengguna", "Cocok jika jumlah Auditor dan PIC tetap di bawah batas"],
    ["Static Assets", "20.000 file per Worker; 25 MiB per file", "Lebih dari cukup untuk frontend aplikasi"],
], widths=[1.25, 2.55, 3.35])
add_callout(doc, "Arti pembuatan gratis", "Aplikasi dapat berjalan tanpa biaya selama penggunaan tetap di bawah kuota gratis. Domain kustom mungkin berbayar jika belum dimiliki. Untuk pilot dapat menggunakan alamat workers.dev. Free plan tidak menyediakan SLA produksi dan layanan dapat berhenti sementara ketika batas harian tercapai.", fill=CREAM, accent=NAVY)

add_heading(doc, "3.5 Tahapan dari awal sampai selesai", 2)
add_table(doc, ["Tahap", "Hasil yang harus selesai", "Gate"], [
    ["0 Validasi aturan", "Rumus skor, skor tiap opsi, bobot, kategori nilai, kewajiban bukti", "Disetujui pemilik proses QAS"],
    ["1 Fondasi", "Repository, React PWA, Worker API, lingkungan dev, D1 dan R2", "Deployment dev berhasil"],
    ["2 Identitas dan master", "Access, RBAC, tiga depo, pengguna, template audit berversi", "Matriks izin lulus uji"],
    ["3 Self audit", "Draft, pilihan jawaban, catatan, bukti, validasi, submit dan lock", "PIC pilot dapat menyelesaikan audit"],
    ["4 Audit resmi", "Mode blind, input Auditor, finalisasi", "Auditor tidak melihat self answer sebelum submit"],
    ["5 Perbandingan dan dashboard", "Gap, tren, heatmap, status, ekspor", "Angka cocok dengan perhitungan uji"],
    ["6 Pilot", "Karawang lalu Baros dan Cirebon", "UAT dan perbaikan selesai"],
    ["7 Produksi", "Backup, monitoring kuota, SOP pengguna, peluncuran", "Owner operasional ditetapkan"],
], widths=[1.05, 4.35, 1.75])

doc.add_page_break()
add_heading(doc, "4 Step 3 Product Requirements Document", 1)
add_heading(doc, "4.1 Nama produk", 2)
add_body(doc, "QAS Audit Hub")
add_heading(doc, "4.2 Pernyataan masalah", 2)
add_body(doc, "Proses audit yang bergantung pada form file menyulitkan input lapangan, pembandingan self audit dengan audit resmi, pemantauan status lintas depo, penyimpanan bukti, dan analisis tren. Sistem perlu memastikan kedua audit menggunakan versi pertanyaan dan aturan nilai yang sama tanpa menghilangkan independensi Auditor.")
add_heading(doc, "4.3 Sasaran produk", 2)
add_bullets(doc, [
    "Mendigitalkan seluruh pertanyaan dan pilihan jawaban pada Form Audit QAS.",
    "Menjamin self audit selesai sebelum audit resmi dimulai.",
    "Menjaga audit resmi independen melalui mode blind.",
    "Menghasilkan perbandingan yang dapat ditelusuri hingga jawaban dan bukti.",
    "Memberikan dashboard depo, periode, bagian, pertanyaan, dan status.",
    "Menjaga biaya awal pada tingkat gratis Cloudflare.",
])
add_heading(doc, "4.4 Indikator keberhasilan", 2)
add_table(doc, ["Indikator", "Target MVP"], [
    ["Kelengkapan", "100 persen pertanyaan wajib dijawab sebelum submit"],
    ["Keterlacakan", "Setiap submit, reopen, dan finalisasi tercatat dengan pengguna dan waktu"],
    ["Perbandingan", "Gap tersedia otomatis setelah audit resmi disubmit"],
    ["Perangkat", "Alur utama dapat diselesaikan pada lebar 360 px dan desktop 1366 px"],
    ["Bukti", "Unggah foto dari kamera ponsel dan lihat ulang melalui izin yang benar"],
    ["Keamanan", "PIC tidak dapat mengakses data depo lain; Auditor dapat mengakses depo yang ditugaskan"],
    ["Kesesuaian skor", "Hasil tes cocok dengan contoh dan rumus yang disetujui owner QAS"],
], widths=[1.7, 5.45])

add_heading(doc, "4.5 Di luar cakupan MVP", 2)
add_bullets(doc, [
    "Aplikasi native Android atau iOS terpisah",
    "Integrasi WhatsApp Business API",
    "Tanda tangan elektronik tersertifikasi",
    "Corrective action workflow lengkap dengan SLA dan eskalasi",
    "Integrasi otomatis dengan sistem HR, dealer, atau ERP",
    "Analisis berbasis AI terhadap foto atau narasi audit",
])

add_heading(doc, "5 Pengguna Peran dan Hak Akses", 1)
add_table(doc, ["Kemampuan", "PIC QAS Gudang", "Auditor QAS"], [
    ["Login", "Ya", "Ya"],
    ["Akses depo", "Hanya depo yang ditetapkan", "Semua atau depo yang ditugaskan"],
    ["Buat dan edit self audit", "Ya sebelum submit", "Tidak"],
    ["Submit self audit", "Ya", "Tidak"],
    ["Lihat jawaban self audit", "Milik depo sendiri", "Setelah audit resmi disubmit secara default"],
    ["Buat dan edit audit resmi", "Tidak", "Ya"],
    ["Submit dan finalisasi audit resmi", "Tidak", "Ya"],
    ["Lihat perbandingan", "Depo sendiri setelah dibuka", "Ya"],
    ["Kelola periode dan template", "Tidak", "Ya, melalui kapabilitas admin"],
    ["Kelola pengguna dan penugasan", "Tidak", "Ya, melalui kapabilitas admin"],
    ["Reopen audit", "Tidak", "Ya dengan alasan"],
    ["Ekspor", "Depo sendiri", "Semua akses yang diizinkan"],
], widths=[3.25, 1.95, 1.95])
add_body(doc, "Tidak diperlukan peran login ketiga pada MVP. Kapabilitas administrasi diberikan kepada Auditor tertentu melalui flag can_manage_master. Hal ini mempertahankan dua tipe login yang diminta sekaligus menyediakan pengelolaan master.")

add_heading(doc, "6 Alur Proses Audit", 1)
add_table(doc, ["Urutan", "Status", "Aktor", "Aturan"], [
    ["1", "PLANNED", "Auditor", "Periode, depo, template, dan jadwal ditetapkan"],
    ["2", "SELF_DRAFT", "PIC", "Jawaban dapat disimpan bertahap dan bukti dapat ditambahkan"],
    ["3", "SELF_SUBMITTED", "PIC", "Semua validasi lulus; self audit terkunci"],
    ["4", "OFFICIAL_DRAFT", "Auditor", "Audit resmi dibuka; jawaban self audit tetap tersembunyi"],
    ["5", "OFFICIAL_SUBMITTED", "Auditor", "Audit resmi terkunci dan perbandingan dihitung"],
    ["6", "COMPARISON_READY", "Sistem", "Gap per pertanyaan, bagian, dan total tersedia"],
    ["7", "FINALIZED", "Auditor", "Hasil final; PIC dapat melihat dan melakukan acknowledgement"],
    ["Khusus", "REOPENED", "Auditor admin", "Alasan wajib; versi dan histori tetap tersimpan"],
], widths=[0.6, 1.5, 1.25, 3.8])
add_heading(doc, "Aturan transisi", 2)
add_bullets(doc, [
    "Audit resmi tidak dapat dimulai sebelum self audit berstatus SELF_SUBMITTED, kecuali override Auditor admin dengan alasan.",
    "Self audit dan audit resmi harus mengacu ke audit_template_version_id yang sama.",
    "Audit yang sudah disubmit tidak dapat diedit biasa.",
    "Reopen membuat event audit dan mempertahankan snapshot sebelum perubahan.",
    "Perbandingan diperbarui hanya dari versi submit terakhir yang sah.",
])

add_heading(doc, "7 Kebutuhan Fungsional", 1)
requirements = [
    ["AUTH-01", "Sistem menerima identitas terverifikasi dari Cloudflare Access."],
    ["AUTH-02", "Sistem mencocokkan email atau subject Access dengan pengguna aktif di D1."],
    ["AUTH-03", "Setiap endpoint memeriksa peran, depo, status audit, dan kepemilikan."],
    ["MST-01", "Auditor admin dapat mengelola Karawang, Baros, dan Cirebon."],
    ["MST-02", "Auditor admin dapat mengelola pengguna, peran, depo, dan status aktif."],
    ["TPL-01", "Template audit memiliki versi, tanggal berlaku, pertanyaan, pilihan, skor, bobot, urutan, dan aturan bukti."],
    ["TPL-02", "Versi yang sudah dipakai audit tidak dapat diubah; perubahan membuat versi baru."],
    ["TPL-03", "Auditor dengan kapabilitas admin dapat menambah, mengedit, mengurutkan, menonaktifkan, dan menghapus pertanyaan pada template berstatus DRAFT."],
    ["TPL-04", "Setiap pertanyaan memiliki kode unik, bagian, redaksi, urutan, bobot, status wajib, aturan bukti, izin N/A, dan status aktif."],
    ["TPL-05", "Setiap pertanyaan memiliki daftar pilihan jawaban yang dapat ditambah, diedit, diurutkan, dinonaktifkan, dan dihapus selama template masih DRAFT."],
    ["TPL-06", "Setiap pilihan jawaban memiliki kode, label, numeric_score, urutan, penanda N/A, dan status aktif."],
    ["TPL-07", "Template PUBLISHED bersifat immutable. Perubahan soal, pilihan, nilai, bobot, atau aturan dilakukan dengan clone ke versi DRAFT baru."],
    ["TPL-08", "Pertanyaan atau pilihan yang sudah pernah digunakan tidak dapat dihapus permanen; item dinonaktifkan atau dihentikan melalui versi baru."],
    ["TPL-09", "Sebelum publish, sistem menyediakan preview form dan simulasi skor serta menolak data yang tidak lengkap atau tidak valid."],
    ["TPL-10", "Audit menyimpan snapshot redaksi pertanyaan, label pilihan, nilai pilihan, bobot, dan konfigurasi skor agar histori tidak berubah."],
    ["CYC-01", "Auditor dapat membuat periode audit per depo dan menetapkan tanggal self audit serta audit resmi."],
    ["AUD-01", "PIC dapat membuat, menyimpan, melanjutkan, dan menyerahkan self audit."],
    ["AUD-02", "Auditor dapat membuat, menyimpan, melanjutkan, dan menyerahkan audit resmi."],
    ["AUD-03", "Setiap jawaban mendukung satu pilihan, skor hasil konfigurasi, dan catatan dokumentasi."],
    ["AUD-04", "Sistem menampilkan progres jumlah pertanyaan selesai dan kekurangan sebelum submit."],
    ["AUD-05", "Submit membutuhkan konfirmasi dan menghasilkan timestamp server."],
    ["EVD-01", "Pengguna dapat mengambil foto dari kamera atau memilih file."],
    ["EVD-02", "Sistem mengompres foto di klien, menampilkan progres upload, dan menyimpan metadata."],
    ["EVD-03", "Bukti tersimpan privat dan hanya dapat dibuka melalui pemeriksaan izin."],
    ["CMP-01", "Sistem menghitung selisih skor official dikurangi self per pertanyaan."],
    ["CMP-02", "Sistem menandai jawaban berbeda, gap negatif, gap positif, dan gap nol."],
    ["CMP-03", "Sistem menampilkan bukti dan catatan kedua audit secara berdampingan setelah dibuka."],
    ["DSH-01", "Dashboard mendukung filter periode, depo, jenis audit, status, bagian, dan kategori."],
    ["DSH-02", "Dashboard menampilkan skor self, skor official, gap, progres, tren, dan pertanyaan dengan gap terbesar."],
    ["EXP-01", "Pengguna dapat mengekspor ringkasan dan detail ke CSV; PDF ringkasan menjadi target fase berikutnya jika diperlukan."],
    ["LOG-01", "Sistem mencatat create, edit penting, submit, reopen, finalisasi, upload, dan delete bukti."],
    ["OFF-01", "PWA menyimpan draft lokal dan antrean sinkronisasi saat koneksi terganggu."],
    ["OFF-02", "Jika konflik versi terjadi, sistem tidak menimpa data server secara diam-diam."],
]
add_table(doc, ["ID", "Requirement"], requirements, widths=[1.0, 6.15])

add_heading(doc, "7.1 Master Template Soal dan Pilihan", 2)
add_body(doc, "Master Soal menjadi sumber form audit. Pengelolaan dilakukan di dalam versi template berstatus DRAFT. Admin dapat mengubah struktur tanpa deployment kode, tetapi perubahan yang sudah diterbitkan tidak boleh mengubah audit historis.")
add_table(doc, ["Objek", "Field minimum", "Operasi pada DRAFT"], [
    ["Template", "kode, nama, versi, tanggal berlaku, status", "Buat, edit metadata, clone, publish, retire"],
    ["Bagian", "kode, nama, urutan, bobot, status aktif", "Tambah, edit, urutkan, hapus jika belum digunakan"],
    ["Soal", "kode, redaksi, bagian, urutan, bobot, wajib, bukti wajib, izin N/A, status aktif", "Tambah, edit, duplikasi, urutkan, nonaktifkan, hapus"],
    ["Pilihan", "kode, label, numeric_score, urutan, is_na, status aktif", "Tambah, edit, duplikasi, urutkan, nonaktifkan, hapus"],
], widths=[1.25, 3.65, 2.25], font_size=8.0)

add_heading(doc, "7.2 Aturan tambah edit hapus", 2)
add_table(doc, ["Kondisi", "Tambah", "Edit", "Hapus"], [
    ["Template DRAFT dan belum digunakan", "Diizinkan", "Diizinkan", "Diizinkan jika tidak merusak referensi draft"],
    ["Template PUBLISHED", "Tidak langsung; clone ke versi baru", "Tidak langsung; clone ke versi baru", "Tidak diizinkan"],
    ["Soal atau pilihan pernah digunakan", "Melalui versi baru", "Melalui versi baru", "Soft delete atau inactive pada versi baru"],
    ["Template RETIRED", "Tidak", "Tidak", "Tidak; hanya untuk histori"],
], widths=[2.15, 1.55, 1.75, 1.7], font_size=7.9)

add_heading(doc, "7.3 Validasi sebelum publish", 2)
add_bullets(doc, [
    "Kode bagian, kode soal, dan kode pilihan harus unik dalam ruang lingkupnya.",
    "Setiap soal aktif memiliki sedikitnya dua pilihan aktif, kecuali tipe input lain disetujui pada versi berikutnya.",
    "Pilihan aktif wajib memiliki label dan numeric_score; aturan N/A harus eksplisit.",
    "Urutan tidak boleh duplikat dalam bagian atau daftar pilihan yang sama.",
    "Bobot, rumus agregasi, pembulatan, dan batas kategori harus lengkap sebelum publish.",
    "Preview harus menampilkan urutan form mobile dan desktop serta simulasi skor maksimum, minimum, dan kasus N/A.",
    "Publish harus mencatat actor, waktu, versi, dan hash konfigurasi ke audit trail.",
])

add_heading(doc, "7.4 Alur perubahan master", 2)
add_numbered(doc, [
    "Admin memilih template PUBLISHED terakhir dan menjalankan Clone as New Draft.",
    "Admin menambah, mengedit, mengurutkan, menonaktifkan, atau menghapus soal dan pilihan pada draft.",
    "Admin mengatur numeric_score pada setiap pilihan dan bobot atau aturan bukti pada setiap soal.",
    "Sistem menjalankan validasi, preview, simulasi skor, dan menampilkan perubahan dibanding versi sebelumnya.",
    "Admin memublikasikan versi setelah keputusan scoring disahkan. Siklus baru dapat memilih versi tersebut; audit lama tetap memakai versi sebelumnya.",
])

add_heading(doc, "8 Struktur Halaman dan Pengalaman Pengguna", 1)
add_table(doc, ["Halaman", "Mobile", "Desktop"], [
    ["Login", "Redirect ke Cloudflare Access dan kembali ke aplikasi", "Sama"],
    ["Beranda PIC", "Audit aktif, progres, due date, tombol lanjutkan", "Kartu status dan histori depo"],
    ["Beranda Auditor", "Daftar tugas dan status per depo", "Dashboard lintas depo dan tabel audit"],
    ["Form audit", "Satu pertanyaan per kartu, navigasi berikutnya, tombol simpan sticky", "Daftar pertanyaan dan panel detail dua kolom"],
    ["Bukti", "Kamera, galeri, kompres, preview, retry", "Drag and drop, preview, download"],
    ["Review submit", "Daftar kekurangan dan ringkasan skor", "Ringkasan bagian dan validasi"],
    ["Comparison", "Filter gap dan kartu per pertanyaan", "Tabel berdampingan self dan official"],
    ["Dashboard", "Kartu KPI dan grafik vertikal", "KPI, grafik, heatmap, dan tabel lengkap"],
    ["Master", "Lihat template dan preview; edit utama disarankan pada desktop", "Pengguna, template, versi, bagian, soal, pilihan, nilai, periode, preview, simulasi, dan publish"],
], widths=[1.25, 2.75, 3.15])

add_heading(doc, "Pola navigasi", 2)
add_bullets(doc, [
    "Mobile menggunakan bottom navigation: Beranda, Audit, Hasil, Profil.",
    "Desktop menggunakan sidebar: Dashboard, Audit, Comparison, Master, Export.",
    "Form menampilkan indikator bagian, progres, status penyimpanan, dan jumlah bukti.",
    "Tombol utama memiliki target sentuh minimal 44 piksel dan teks dapat dibaca tanpa zoom.",
    "Warna tidak menjadi satu-satunya penanda gap atau status; selalu sertakan label dan ikon.",
])

add_heading(doc, "9 Dashboard dan Perbandingan", 1)
add_heading(doc, "9.1 Kartu ringkasan", 2)
add_bullets(doc, [
    "Jumlah audit planned, self submitted, official submitted, dan finalized",
    "Nilai rata-rata self audit",
    "Nilai rata-rata audit resmi",
    "Gap rata-rata official minus self",
    "Kelengkapan bukti",
    "Depo dengan gap terbesar",
])
add_heading(doc, "9.2 Visual utama", 2)
add_table(doc, ["Visual", "Tujuan"], [
    ["Bar chart per depo", "Bandingkan self dan official Karawang, Baros, dan Cirebon"],
    ["Trend per periode", "Melihat perubahan skor official dan gap dari waktu ke waktu"],
    ["Heatmap pertanyaan", "Menemukan pertanyaan yang sering memiliki gap negatif"],
    ["Breakdown per bagian", "Bandingkan J1 distribusi, NRFS, maintenance, dan J2 distribusi"],
    ["Status funnel", "Melihat audit yang tertahan sebelum self submit, official submit, atau finalisasi"],
    ["Tabel gap", "Drill down ke jawaban, catatan, dan bukti"],
], widths=[2.0, 5.15])
add_heading(doc, "9.3 Definisi gap", 2)
add_body(doc, "Gunakan gap = skor audit resmi - skor self audit. Gap negatif berarti self assessment lebih tinggi daripada hasil Auditor. Gap positif berarti hasil Auditor lebih tinggi. Dashboard harus menampilkan tanda, angka, dan label agar interpretasi tidak bergantung pada warna.")

add_heading(doc, "10 Aturan Penilaian", 1)
add_body(doc, "Skor tidak boleh dikodekan langsung di frontend. Setiap pilihan jawaban memiliki numeric_score dan setiap pertanyaan dapat memiliki weight. Template versi menyimpan formula agregasi dan rentang kategori. Audit menyimpan snapshot hasil per pertanyaan dan total pada saat submit agar histori tidak berubah ketika template baru dibuat.")
add_heading(doc, "Model perhitungan yang direkomendasikan", 2)
add_table(doc, ["Komponen", "Aturan"], [
    ["Skor jawaban", "Nilai dari answer_option.numeric_score"],
    ["Skor bagian", "Jumlah skor berbobot dibagi jumlah bobot yang berlaku"],
    ["Nilai akhir", "Formula konfigurasi template; default belum ditetapkan sampai owner QAS mengonfirmasi"],
    ["Kategori", "Rentang min dan max berversi; contoh PDF menunjukkan 4,37 sebagai BAIK SEKALI"],
    ["Tidak berlaku", "Hanya tersedia jika disetujui dan harus memiliki aturan denominator yang jelas"],
    ["Pembulatan", "Jumlah desimal dan metode pembulatan disimpan pada template"],
], widths=[1.7, 5.45])
add_callout(doc, "Temuan penting dari form", "Jika angka Point yang terlihat dijumlahkan lalu dibagi 17 pertanyaan, hasilnya tidak persis sama dengan nilai akhir 4,37 pada PDF. Ini menunjukkan adanya aturan, bobot, skor pilihan, atau presisi yang tidak terlihat. Build tidak boleh dimulai pada mesin skor final sebelum aturan sumber dikonfirmasi.", fill="FCE4D6", accent=RED)

add_heading(doc, "11 Arsitektur dan Penyimpanan Data", 1)
add_table(doc, ["Klien", "Edge", "Data", "Identitas", "Operasional"], [
    ["React PWA\nIndexedDB draft", "Cloudflare Worker\nStatic Assets dan API", "D1 SQL\nR2 private evidence", "Cloudflare Access\nOTP atau Cloudflare IdP", "Wrangler\nGit\nLogs dan Analytics"],
], widths=[1.45, 1.55, 1.35, 1.55, 1.25], font_size=8.6)
add_heading(doc, "Alur request", 2)
add_numbered(doc, [
    "Pengguna membuka aplikasi dan melewati autentikasi Cloudflare Access.",
    "Worker memvalidasi JWT Access dan mengambil profil aplikasi dari D1.",
    "API memeriksa role dan depot_scope sebelum membaca atau menulis data.",
    "Jawaban disimpan ke D1 melalui transaksi atau batch yang terbatas.",
    "Bukti diunggah ke endpoint yang memeriksa izin, kemudian disimpan di R2 private bucket.",
    "Dashboard menjalankan query terindeks atau membaca tabel agregat untuk menghemat row reads.",
])
add_heading(doc, "Strategi lingkungan", 2)
add_table(doc, ["Lingkungan", "Database", "Bucket", "URL"], [
    ["Local", "D1 local", "R2 mock atau dev", "localhost"],
    ["Preview", "qas-dev", "qas-evidence-dev", "preview worker"],
    ["Production", "qas-prod", "qas-evidence-prod", "workers.dev atau domain perusahaan"],
], widths=[1.3, 1.8, 2.0, 2.05])

add_heading(doc, "12 Model Data", 1)
data_tables = [
    ["depots", "id, code, name, active", "Karawang, Baros, Cirebon"],
    ["users", "id, access_sub, email, name, active", "Identitas aplikasi"],
    ["user_roles", "user_id, role, depot_id, can_manage_master", "RBAC dan cakupan depo"],
    ["audit_templates", "id, code, name", "Identitas template"],
    ["audit_template_versions", "id, template_id, version, valid_from, scoring_config_json, status", "Versi immutable"],
    ["sections", "id, template_version_id, code, name, sort_order, weight, active", "J1 dan J2 serta subsection"],
    ["questions", "id, section_id, code, text, weight, required, evidence_rule, allow_na, sort_order, active", "Master pertanyaan audit"],
    ["answer_options", "id, question_id, code, text, numeric_score, is_na, sort_order, active", "Master pilihan dan nilai"],
    ["audit_cycles", "id, period, depot_id, template_version_id, dates, status", "Pasangan self dan official"],
    ["audits", "id, cycle_id, type, status, auditor_user_id, score, category, timestamps, version", "SELF atau OFFICIAL"],
    ["audit_answers", "id, audit_id, question_id, option_id, question_text_snapshot, option_text_snapshot, score_snapshot, weight_snapshot, notes", "Jawaban dan snapshot historis"],
    ["evidence_files", "id, answer_id, r2_key, filename, mime_type, size, hash, uploaded_by", "Metadata bukti"],
    ["audit_events", "id, audit_id, event_type, actor_id, at, reason, payload_json", "Audit trail"],
    ["comparison_snapshots", "id, cycle_id, calculated_at, self_score, official_score, gap, data_json", "Hasil final yang dapat diaudit"],
    ["acknowledgements", "id, cycle_id, user_id, at, note", "PIC mengakui hasil"],
]
add_table(doc, ["Tabel", "Kolom utama", "Fungsi"], data_tables, widths=[1.75, 3.5, 1.9], font_size=7.8)
add_heading(doc, "Indeks minimum", 2)
add_bullets(doc, [
    "users(email) dan users(access_sub)",
    "user_roles(user_id, depot_id)",
    "audit_cycles(depot_id, period, status)",
    "audits(cycle_id, type) dengan unique constraint untuk satu self dan satu official",
    "audit_answers(audit_id, question_id) dengan unique constraint",
    "questions(template_version_id melalui section, sort_order)",
    "audit_events(audit_id, at)",
])

add_heading(doc, "13 Kontrak API", 1)
api_rows = [
    ["GET", "/api/me", "Profil, role, dan depot scope"],
    ["GET", "/api/depots", "Daftar depo yang dapat diakses"],
    ["GET", "/api/cycles", "Daftar periode dan status dengan filter"],
    ["POST", "/api/cycles", "Buat periode audit"],
    ["GET", "/api/cycles/:id", "Detail cycle dan ringkasan"],
    ["POST", "/api/cycles/:id/self", "Buat self audit"],
    ["POST", "/api/cycles/:id/official", "Buat audit resmi setelah syarat terpenuhi"],
    ["GET", "/api/audits/:id", "Detail audit sesuai izin dan blind rule"],
    ["PUT", "/api/audits/:id/answers/:questionId", "Simpan jawaban dengan optimistic version"],
    ["POST", "/api/audits/:id/submit", "Validasi, hitung, snapshot, dan lock"],
    ["POST", "/api/audits/:id/reopen", "Buka kembali dengan alasan"],
    ["POST", "/api/answers/:id/evidence", "Unggah bukti"],
    ["GET", "/api/evidence/:id", "Akses bukti setelah pemeriksaan izin"],
    ["DELETE", "/api/evidence/:id", "Hapus bukti sebelum submit atau melalui reopen"],
    ["GET", "/api/cycles/:id/comparison", "Perbandingan self dan official"],
    ["POST", "/api/cycles/:id/finalize", "Finalisasi oleh Auditor"],
    ["GET", "/api/dashboard", "Agregasi dengan filter"],
    ["GET", "/api/exports/audits.csv", "Ekspor data sesuai scope"],
    ["GET", "/api/admin/templates", "Daftar template dan versi"],
    ["POST", "/api/admin/templates/:id/versions", "Buat versi baru"],
    ["POST", "/api/admin/template-versions/:id/clone", "Clone versi menjadi DRAFT baru"],
    ["GET", "/api/admin/template-versions/:id", "Detail bagian, soal, pilihan, dan konfigurasi"],
    ["POST", "/api/admin/template-versions/:id/sections", "Tambah bagian pada DRAFT"],
    ["PATCH", "/api/admin/sections/:id", "Edit atau ubah urutan bagian DRAFT"],
    ["DELETE", "/api/admin/sections/:id", "Hapus bagian yang aman pada DRAFT"],
    ["POST", "/api/admin/sections/:id/questions", "Tambah soal pada DRAFT"],
    ["PATCH", "/api/admin/questions/:id", "Edit, urutkan, atau nonaktifkan soal DRAFT"],
    ["DELETE", "/api/admin/questions/:id", "Hapus soal yang aman pada DRAFT"],
    ["POST", "/api/admin/questions/:id/options", "Tambah pilihan dan numeric_score"],
    ["PATCH", "/api/admin/options/:id", "Edit label, nilai, N/A, urutan, atau status"],
    ["DELETE", "/api/admin/options/:id", "Hapus pilihan yang aman pada DRAFT"],
    ["POST", "/api/admin/template-versions/:id/validate", "Validasi dan simulasi scoring"],
    ["POST", "/api/admin/template-versions/:id/publish", "Publish versi immutable"],
]
add_table(doc, ["Method", "Endpoint", "Tujuan"], api_rows, widths=[0.75, 2.9, 3.5], font_size=7.9)

add_heading(doc, "14 Keamanan Privasi dan Audit Trail", 1)
add_bullets(doc, [
    "Validasi signature, issuer, audience, dan expiry JWT Cloudflare Access pada Worker.",
    "Gunakan deny by default untuk semua endpoint dan cek depot_scope pada server.",
    "Gunakan satu origin untuk frontend dan API agar risiko CORS dan cookie lebih sederhana.",
    "R2 bucket bersifat privat; jangan gunakan public bucket untuk bukti audit.",
    "Nama object R2 menggunakan ID acak, bukan email, nama pengguna, atau nama file asli.",
    "Batasi jenis file, ukuran, jumlah, dan lakukan verifikasi MIME pada server.",
    "Simpan waktu dalam UTC dan tampilkan Asia Jakarta pada UI.",
    "Tidak ada hard delete untuk audit finalized. Koreksi dilakukan melalui versi atau reopen yang terlacak.",
    "Jangan mencatat isi sensitif, JWT, atau URL bukti sementara ke log aplikasi.",
    "Tambahkan rate limiting untuk upload dan endpoint sensitif.",
])
add_heading(doc, "Audit event minimum", 2)
add_table(doc, ["Event", "Data minimum"], [
    ["AUDIT_CREATED", "audit, actor, waktu, template version"],
    ["ANSWER_CHANGED", "audit, question, actor, waktu, old hash, new hash"],
    ["EVIDENCE_UPLOADED", "answer, file metadata, actor, waktu"],
    ["AUDIT_SUBMITTED", "audit, skor, kategori, actor, waktu"],
    ["AUDIT_REOPENED", "audit, actor, waktu, alasan"],
    ["AUDIT_FINALIZED", "cycle, comparison snapshot, actor, waktu"],
    ["EXPORT_CREATED", "filter, actor, waktu"],
    ["TEMPLATE_VERSION_CLONED", "template, source version, new draft, actor, waktu"],
    ["TEMPLATE_MASTER_CHANGED", "version, object type, object id, old hash, new hash, actor, waktu"],
    ["TEMPLATE_VERSION_PUBLISHED", "template, version, configuration hash, actor, waktu"],
], widths=[2.1, 5.05])

add_heading(doc, "15 Keputusan yang Wajib Dikonfirmasi", 1)
decisions = [
    ["D-01", "Skor numerik setiap pilihan jawaban", "Owner QAS", "Blocker mesin penilaian"],
    ["D-02", "Rumus nilai akhir, bobot per pertanyaan atau bagian, dan pembulatan", "Owner QAS", "Blocker mesin penilaian"],
    ["D-03", "Rentang kategori termasuk BAIK SEKALI", "Owner QAS", "Blocker dashboard kategori"],
    ["D-04", "Pertanyaan yang wajib dan aturan bukti per pertanyaan", "Owner QAS", "Blocker validasi submit"],
    ["D-05", "Apakah Auditor harus blind terhadap self audit sampai official submit", "Pemilik proses", "Direkomendasikan Ya"],
    ["D-06", "Daftar email, nama, role, dan depo pengguna", "Admin QAS", "Diperlukan sebelum UAT"],
    ["D-07", "Frekuensi periode audit dan due date", "Pemilik proses", "Diperlukan untuk cycle"],
    ["D-08", "Retensi bukti dan kebijakan penghapusan", "Pemilik data", "Diperlukan sebelum produksi"],
    ["D-09", "Format ekspor resmi yang dibutuhkan", "Auditor QAS", "CSV ada di MVP; PDF perlu format"],
    ["D-10", "Daftar Auditor yang memperoleh can_manage_master dan kewenangan publish", "Owner QAS", "Kontrol perubahan Master Soal"],
]
add_table(doc, ["ID", "Keputusan", "Pemilik", "Dampak"], decisions, widths=[0.7, 3.4, 1.35, 1.7], font_size=8)

add_heading(doc, "16 Acceptance Criteria", 1)
criteria = [
    ["AC-01", "PIC Karawang tidak dapat melihat audit Baros atau Cirebon melalui UI maupun API."],
    ["AC-02", "Auditor dapat melihat audit sesuai depot assignment dan Auditor admin dapat mengelola master."],
    ["AC-03", "Self audit tidak dapat disubmit jika pertanyaan wajib atau bukti wajib belum lengkap."],
    ["AC-04", "Setelah self submit, edit biasa ditolak oleh server."],
    ["AC-05", "Audit resmi tidak dapat dibuat sebelum self submit kecuali override beralasan."],
    ["AC-06", "Dalam mode blind, endpoint audit resmi tidak mengembalikan jawaban self sebelum official submit."],
    ["AC-07", "Setelah official submit, comparison menampilkan opsi, skor, gap, catatan, dan bukti kedua audit."],
    ["AC-08", "Nilai hasil uji cocok dengan perhitungan independen berdasarkan konfigurasi template."],
    ["AC-09", "Foto dari kamera ponsel dapat diunggah, dilihat ulang, dan tetap privat."],
    ["AC-10", "Draft offline dapat dipulihkan dan disinkronkan; konflik tidak menimpa data server tanpa peringatan."],
    ["AC-11", "Dashboard filter depo dan periode menghasilkan angka konsisten dengan detail audit."],
    ["AC-12", "Reopen membutuhkan alasan dan menghasilkan event audit."],
    ["AC-13", "Template versi lama tetap menghasilkan tampilan historis yang sama setelah versi baru dibuat."],
    ["AC-14", "Aplikasi dapat digunakan tanpa horizontal scroll pada lebar 360 px untuk alur audit utama."],
    ["AC-15", "Ekspor hanya mencakup data yang diizinkan untuk pengguna tersebut."],
    ["AC-16", "Auditor admin dapat menambah, mengedit, mengurutkan, menonaktifkan, dan menghapus soal serta pilihan pada versi DRAFT."],
    ["AC-17", "Nilai setiap pilihan dapat diatur per soal dan dipakai oleh simulasi serta scoring engine dari sumber konfigurasi yang sama."],
    ["AC-18", "API menolak perubahan langsung pada bagian, soal, pilihan, nilai, bobot, dan aturan versi PUBLISHED atau RETIRED."],
    ["AC-19", "Pertanyaan atau pilihan yang pernah digunakan tidak dapat dihapus permanen dan audit historis tetap menampilkan redaksi serta nilai snapshot."],
    ["AC-20", "Publish ditolak jika ada kode atau urutan duplikat, soal tanpa pilihan aktif, nilai kosong, atau konfigurasi scoring belum lengkap."],
    ["AC-21", "Preview dan simulasi menampilkan hasil maksimum, minimum, N/A, pembulatan, serta kategori sebelum versi diterbitkan."],
]
add_table(doc, ["ID", "Kriteria penerimaan"], criteria, widths=[0.85, 6.3])

add_heading(doc, "17 Rencana Implementasi Antigravity", 1)
add_body(doc, "Gunakan Antigravity sebagai coding agent dengan pekerjaan kecil, acceptance criteria yang jelas, dan verifikasi browser pada setiap milestone. Jangan memberikan satu prompt besar untuk membangun seluruh aplikasi sekaligus.")
add_heading(doc, "17.1 Struktur repository", 2)
p = doc.add_paragraph(style="Code")
p.add_run("""qas-audit-app/
  src/                 React UI
  worker/              Hono API dan middleware
  migrations/          D1 SQL migrations
  seed/                Depo dan template versi awal
  tests/               Unit integration dan e2e
  docs/                PRD data model dan decisions
  public/              PWA icons dan manifest
  wrangler.jsonc       Binding D1 R2 dan environment
  AGENTS.md            Aturan kerja Antigravity
  README.md            Setup dan deployment""")

add_heading(doc, "17.2 Urutan prompt atau task", 2)
tasks = [
    ["Task 1", "Scaffold React TypeScript Vite PWA dan Cloudflare Worker Hono dalam satu repository. Tambahkan health endpoint dan deployment dev."],
    ["Task 2", "Buat migrasi D1 untuk master, template versioning, audit, answers, evidence metadata, dan audit events. Tambahkan seed tiga depo."],
    ["Task 3", "Integrasikan Cloudflare Access JWT validation dan RBAC. Buat test matriks role dan depot."],
    ["Task 4", "Buat Master Template, Bagian, Soal, Pilihan, nilai per pilihan, clone version, validasi, preview, simulasi, publish, dan retire. Terapkan aturan CRUD hanya pada DRAFT."],
    ["Task 5", "Buat template reader dan audit cycle. Jangan implementasikan skor final sebelum keputusan D-01 sampai D-04 diisi."],
    ["Task 6", "Buat self audit mobile first dengan autosave, validation, evidence upload, review, submit, dan lock."],
    ["Task 7", "Buat official audit dengan blind rule yang diuji pada API dan UI."],
    ["Task 8", "Buat scoring engine konfigurabel dan comparison snapshot. Tambahkan golden test dari contoh yang disetujui."],
    ["Task 9", "Buat dashboard, filter, drill down, dan CSV export."],
    ["Task 10", "Tambahkan offline draft, conflict handling, accessibility, error handling, dan performance checks."],
    ["Task 11", "Jalankan UAT mobile dan desktop, audit security, verifikasi kuota, lalu deploy production."],
]
add_table(doc, ["Urutan", "Instruksi ke Antigravity"], tasks, widths=[0.9, 6.25], font_size=8.3)

add_heading(doc, "17.3 Guardrails untuk AGENTS.md", 2)
add_bullets(doc, [
    "Gunakan TypeScript strict dan hindari any tanpa alasan.",
    "Semua otorisasi dilakukan kembali di server; UI hiding bukan kontrol keamanan.",
    "Jangan mengubah migration yang sudah diterapkan; buat migration baru.",
    "Semua audit menggunakan template version immutable.",
    "Tidak ada public R2 URL untuk bukti.",
    "Tambahkan test sebelum menutup task untuk scoring, RBAC, state transition, dan blind rule.",
    "Gunakan browser verification pada viewport mobile dan desktop.",
    "Jangan mengarang skor atau kategori yang belum dikonfirmasi.",
])

add_heading(doc, "18 Pengujian Peluncuran dan Operasional", 1)
add_heading(doc, "18.1 Test minimum", 2)
add_table(doc, ["Jenis", "Cakupan"], [
    ["Unit", "Scoring, pembulatan, kategori, state transition, file validation"],
    ["Integration", "D1 transactions, RBAC, Access identity, R2 upload dan delete"],
    ["E2E", "PIC self audit, Auditor official audit, blind comparison, reopen, export"],
    ["Responsive", "360 x 800, 390 x 844, tablet, dan 1366 x 768"],
    ["Offline", "Putus koneksi saat edit dan upload; recovery serta conflict"],
    ["Security", "IDOR lintas depo, file access, JWT invalid, role escalation, upload berbahaya"],
    ["Performance", "Dashboard dan daftar audit menggunakan indeks dan pagination"],
], widths=[1.45, 5.7])

add_heading(doc, "18.2 Pilot dan rollout", 2)
add_numbered(doc, [
    "Masukkan template versi awal dan satu akun Auditor serta satu PIC Karawang.",
    "Jalankan satu self audit dan audit resmi contoh di Karawang.",
    "Rekonsiliasi seluruh skor terhadap perhitungan manual pemilik QAS.",
    "Perbaiki istilah, urutan pertanyaan, validasi bukti, dan tampilan lapangan.",
    "Aktifkan Baros dan Cirebon setelah pilot Karawang diterima.",
    "Tetapkan owner user master, template, backup, dan monitoring kuota.",
])

add_heading(doc, "18.3 Operasional gratis", 2)
add_bullets(doc, [
    "Pantau request Workers, D1 row reads dan writes, storage D1, storage R2, serta operasi R2.",
    "Gunakan query terindeks, pagination, dan agregasi terkontrol untuk dashboard.",
    "Kompres foto pada klien dan gunakan batas awal yang dapat dikonfigurasi, misalnya maksimum 3 foto per pertanyaan dan 2 MB per foto setelah kompresi. Nilai final harus dikonfirmasi owner.",
    "Gunakan D1 Time Travel yang tersedia pada plan untuk pemulihan jangka pendek dan ekspor berkala untuk retensi tambahan.",
    "Siapkan prosedur upgrade apabila penggunaan mendekati batas gratis agar audit tidak terganggu.",
])

add_heading(doc, "19 Risiko dan Mitigasi", 1)
add_table(doc, ["Risiko", "Dampak", "Mitigasi"], [
    ["Aturan skor tidak lengkap", "Hasil tidak sah", "Blokir scoring final sampai D-01 sampai D-04 disetujui; gunakan template konfigurabel"],
    ["Auditor terpengaruh self audit", "Perbandingan bias", "Mode blind pada API dan UI"],
    ["Foto memenuhi kuota R2", "Upload gagal atau biaya muncul", "Kompres, batas ukuran atau jumlah, retensi, monitoring"],
    ["D1 row reads tinggi", "Batas gratis tercapai", "Indeks, pagination, agregasi, hindari full scan"],
    ["PIC mengakses depo lain", "Kebocoran data", "RBAC server side dan test IDOR"],
    ["Koneksi gudang tidak stabil", "Jawaban hilang", "Autosave lokal, antrean sync, status koneksi, conflict handling"],
    ["Template berubah di tengah audit", "Histori tidak konsisten", "Template version immutable"],
    ["Free plan tanpa SLA", "Gangguan operasional", "Jadwal audit, export, backup, dan rencana upgrade"],
    ["Build AI mengarang aturan", "Kesalahan bisnis", "Decision register, golden tests, review manusia per milestone"],
], widths=[1.65, 1.55, 3.95], font_size=8)

add_heading(doc, "20 Sumber", 1)
sources = [
    ("Form Audit QAS.pdf", "F:\\2. QAS\\Form Audit QAS.pdf"),
    ("Cloudflare Workers Pricing", "https://developers.cloudflare.com/workers/platform/pricing/"),
    ("Cloudflare Workers Limits", "https://developers.cloudflare.com/workers/platform/limits/"),
    ("Cloudflare D1 Pricing", "https://developers.cloudflare.com/d1/platform/pricing/"),
    ("Cloudflare D1 Limits", "https://developers.cloudflare.com/d1/platform/limits/"),
    ("Cloudflare R2 Pricing", "https://developers.cloudflare.com/r2/pricing/"),
    ("Cloudflare Access One Time PIN", "https://developers.cloudflare.com/cloudflare-one/integrations/identity-providers/one-time-pin/"),
    ("Cloudflare Access JWT Validation", "https://developers.cloudflare.com/cloudflare-one/access-controls/applications/http-apps/authorization-cookie/validating-json/"),
    ("Cloudflare Zero Trust Plans", "https://www.cloudflare.com/plans/zero-trust-services/"),
    ("Google Antigravity Build with Google", "https://www.antigravity.google/docs/build-with-google"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    if url.startswith("http"):
        add_hyperlink(p, label, url)
    else:
        p.add_run(f"{label} - sumber internal yang dilampirkan")

doc.add_page_break()
add_heading(doc, "Lampiran A Pemetaan Form Audit QAS", 1)
add_body(doc, "Lampiran ini memetakan 17 pertanyaan dan pilihan jawaban yang terbaca dari PDF. Redaksi perlu divalidasi terhadap file Excel sumber sebelum dijadikan template produksi. Kolom Point pada PDF dipertahankan sebagai nilai contoh yang terlihat, bukan sebagai definisi maksimum.")

current_group = None
current_section = None
for q in questions:
    if q["group"] != current_group:
        add_heading(doc, q["group"], 2)
        current_group = q["group"]
        current_section = None
    if q["section"] != current_section:
        add_heading(doc, q["section"], 3)
        current_section = q["section"]
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width = Inches(1.35)
    t.columns[1].width = Inches(5.8)
    rows = [
        ("Kode", q["id"]),
        ("Pertanyaan", q["question"]),
        ("Point terlihat", q["score"]),
        ("Wajib pada PDF", q["mandatory"]),
        ("Pilihan", "\n".join([f"{i+1}. {o}" for i, o in enumerate(q["options"])])),
        ("Input aplikasi", "Pilihan tunggal, catatan dokumentasi, dan bukti file atau foto sesuai konfigurasi template"),
    ]
    for i, (label, value) in enumerate(rows):
        if i > 0:
            t.add_row()
        set_cell_shading(t.cell(i, 0), LIGHT_BLUE)
        set_cell_text(t.cell(i, 0), label, bold=True, color=NAVY, size=8)
        p = set_cell_text(t.cell(i, 1), value, size=8)
        if label == "Pilihan":
            p.paragraph_format.line_spacing = 1.0
    set_table_widths(t, [1.35, 5.8])
    keep_table_together(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

add_heading(doc, "Lampiran B Ringkasan Temuan Form", 1)
add_table(doc, ["Elemen", "Temuan"], [
    ["Jumlah halaman", "2"],
    ["Jumlah pertanyaan", "17"],
    ["Kelompok", "J1 AHM to Main Dealer dan J2 Main Dealer to Dealer"],
    ["Subbagian", "Distribusi AHM ke Main Dealer, NRFS, Maintenance, dan Distribusi Main Dealer ke Dealer"],
    ["Metadata", "Periode, Gudang, Tanggal Audit, Auditor, Auditee"],
    ["Persetujuan", "Dibuat Self Audit, Disetujui Auditee, Diketahui Kepala Gudang"],
    ["Bukti", "Kolom Dokumentasi tersedia untuk setiap pertanyaan"],
    ["Nilai contoh", "4,37"],
    ["Kategori contoh", "BAIK SEKALI"],
    ["Informasi yang belum terlihat", "Skor setiap opsi, bobot, formula total, pembulatan, kategori, dan aturan bukti"],
], widths=[2.0, 5.15])

doc.core_properties.title = "PRD Aplikasi Audit QAS"
doc.core_properties.subject = "Aplikasi self audit dan audit resmi QAS untuk tiga depo"
doc.core_properties.author = "Project QAS"
doc.core_properties.keywords = "QAS, audit, Cloudflare, D1, R2, PWA, Antigravity"

doc.save(OUT)
print(OUT)
