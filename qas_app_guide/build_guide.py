from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\Ari\.codex\.chatgpt-projects\g-p-6aacf1deeee08191b49d027b1363f153")
OUT_DIR = ROOT / "outputs" / "qas_app_guide"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / "Panduan_Lengkap_Pembangunan_Aplikasi_Audit_QAS.docx"

NAVY, NAVY2, BLUE = "0B2D57", "173B67", "2F75B5"
GREEN, CREAM = "70AD47", "F7EBD2"
LIGHT_BLUE, LIGHT_GREEN = "EAF1F8", "E2F0D9"
LIGHT_GREY, AMBER = "F2F4F7", "F4B183"
LIGHT_AMBER, RED, LIGHT_RED = "FFF2CC", "C00000", "FCE4D6"
DARK, WHITE = "17243A", "FFFFFF"


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = pr.find(qn("w:shd")) or OxmlElement("w:shd")
    if el.getparent() is None:
        pr.append(el)
    el.set(qn("w:fill"), fill)


def margins(cell, top=80, start=100, bottom=80, end=100):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def widths(tbl, values):
    for row in tbl.rows:
        for i, value in enumerate(values):
            row.cells[i].width = Inches(value)
            pr = row.cells[i]._tc.get_or_add_tcPr()
            w = pr.find(qn("w:tcW")) or OxmlElement("w:tcW")
            if w.getparent() is None:
                pr.append(w)
            w.set(qn("w:w"), str(int(value * 1440))); w.set(qn("w:type"), "dxa")


def runfmt(run, size=9.3, bold=False, color=DARK, italic=False, font="Arial"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size, run.font.bold, run.font.italic = Pt(size), bold, italic
    run.font.color.rgb = RGBColor.from_string(color)


def celltext(cell, text, bold=False, color=DARK, size=8.2, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None: p.alignment = align
    runfmt(p.add_run(str(text)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)


def page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Halaman "); runfmt(r, 8, color="667085")
    b, ins, e = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin"); ins.set(qn("xml:space"), "preserve"); ins.text = " PAGE "; e.set(qn("w:fldCharType"), "end")
    r._r.extend([b, ins, e])


def hyperlink(p, label, url):
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link, r, pr = OxmlElement("w:hyperlink"), OxmlElement("w:r"), OxmlElement("w:rPr")
    link.set(qn("r:id"), rid)
    fonts, color, under = OxmlElement("w:rFonts"), OxmlElement("w:color"), OxmlElement("w:u")
    fonts.set(qn("w:ascii"), "Arial"); fonts.set(qn("w:hAnsi"), "Arial")
    color.set(qn("w:val"), BLUE); under.set(qn("w:val"), "single")
    pr.extend([fonts, color, under]); r.append(pr)
    t = OxmlElement("w:t"); t.text = label; r.append(t); link.append(r); p._p.append(link)


def configure(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.top_margin, s.bottom_margin = Inches(.62), Inches(.58)
    s.left_margin, s.right_margin = Inches(.67), Inches(.67)
    for name in ["Normal", "Body Text"]:
        st = doc.styles[name]; st.font.name = "Arial"; st.font.size = Pt(9.3); st.font.color.rgb = RGBColor.from_string(DARK)
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial"); st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(5), 1.08
    for level, size, color in [(1, 17, NAVY), (2, 13, NAVY2), (3, 10.5, BLUE)]:
        st = doc.styles[f"Heading {level}"]; st.font.name = "Arial"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial"); st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        st.paragraph_format.space_before, st.paragraph_format.space_after, st.paragraph_format.keep_with_next = Pt(9), Pt(4), True
    for name in ["List Bullet", "List Bullet 2", "List Number"]:
        st = doc.styles[name]; st.font.name = "Arial"; st.font.size = Pt(9.2); st.paragraph_format.space_after = Pt(2.5); st.paragraph_format.line_spacing = 1.05
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial"); st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    p = s.header.paragraphs[0]; runfmt(p.add_run("PANDUAN PEMBANGUNAN APLIKASI AUDIT QAS"), 7.5, True, NAVY)
    p = s.footer.paragraphs[0]; runfmt(p.add_run("Motorcycle Logistic  |  Cloudflare + Antigravity  •  "), 7.5, color="667085"); page_number(p)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.add_run(text); return p


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_prefix and text.startswith(bold_prefix):
        runfmt(p.add_run(bold_prefix), bold=True); runfmt(p.add_run(text[len(bold_prefix):]))
    else: runfmt(p.add_run(text))
    return p


def bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2"); runfmt(p.add_run(item), 9.2)


def numbers(doc, items):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.22)
        p.paragraph_format.space_after = Pt(2.5)
        runfmt(p.add_run(f"{idx}. {item}"), 9.2)


def callout(doc, title, text, fill=LIGHT_BLUE, accent=NAVY):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; widths(t, [.14, 7.0])
    shade(t.cell(0, 0), accent); shade(t.cell(0, 1), fill)
    p = t.cell(0, 1).paragraphs[0]; runfmt(p.add_run(title + "\n"), 9.6, True, accent); runfmt(p.add_run(text), 9.0)
    margins(t.cell(0, 1), 120, 150, 120, 150); doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(doc, headers, rows, col_widths=None, font_size=8.1, header_fill=NAVY):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    trpr = t.rows[0]._tr.get_or_add_trPr(); rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trpr.append(rep)
    for i, h in enumerate(headers): shade(t.rows[0].cells[i], header_fill); celltext(t.rows[0].cells[i], h, True, WHITE, 8.2, WD_ALIGN_PARAGRAPH.CENTER)
    for idx, row in enumerate(rows):
        cells = t.add_row().cells
        if idx % 2:
            for c in cells: shade(c, LIGHT_GREY)
        for i, value in enumerate(row): celltext(cells[i], value, size=font_size)
        cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    if col_widths: widths(t, col_widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def code(doc, text, title=None):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; widths(t, [7.1]); shade(t.cell(0, 0), "111827")
    c = t.cell(0, 0); c.text = ""; p = c.paragraphs[0]
    if title: runfmt(p.add_run(title + "\n"), 8.2, True, AMBER, font="Consolas")
    runfmt(p.add_run(text), 7.6, color="F9FAFB", font="Consolas"); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    margins(c, 110, 130, 110, 130); doc.add_paragraph().paragraph_format.space_after = Pt(0)


def section(doc, number, title, subtitle, pagebreak=True):
    if pagebreak:
        doc.add_page_break()
    p = doc.add_paragraph(); runfmt(p.add_run(f"{number:02d}"), 28, True, GREEN); p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); runfmt(p.add_run(title), 22, True, NAVY); p.paragraph_format.space_after = Pt(5); p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); runfmt(p.add_run(subtitle), 10.5, color="667085"); p.paragraph_format.space_after = Pt(12); p.paragraph_format.keep_with_next = True


def prompt(doc, no, title, text, done):
    heading(doc, f"Prompt {no} — {title}", 2); code(doc, text, "PROMPT UNTUK ANTIGRAVITY"); body(doc, "Kriteria selesai:", "Kriteria selesai:"); bullets(doc, done)


doc = Document(); configure(doc)

# Cover
for _ in range(3): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runfmt(p.add_run("PANDUAN LENGKAP"), 15, True, GREEN)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runfmt(p.add_run("PEMBANGUNAN APLIKASI\nAUDIT QAS"), 30, True, NAVY); p.paragraph_format.space_after = Pt(12)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runfmt(p.add_run("Quality Assurance System • Self Audit • Audit Auditor • Dashboard Perbandingan"), 11, color=NAVY2)
doc.add_paragraph()
t = doc.add_table(rows=1, cols=3); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; widths(t, [2.25, 2.25, 2.25])
for i, (big, small, fill) in enumerate([("3 DEPO", "Karawang • Baros • Cirebon", NAVY), ("2 PERAN", "PIC QAS • Auditor QAS", GREEN), ("1 ALUR", "Self audit → Audit asli → Compare", NAVY2)]):
    shade(t.cell(0, i), fill); p = t.cell(0, i).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runfmt(p.add_run(big + "\n"), 13, True, WHITE); runfmt(p.add_run(small), 8.5, color=WHITE); margins(t.cell(0, i), 200, 100, 200, 100)
doc.add_paragraph(); callout(doc, "ARAH TEKNOLOGI", "React + TypeScript + Vite PWA • Cloudflare Worker + Hono • D1 • R2 • Cloudflare Access • IndexedDB", CREAM, NAVY)
doc.add_paragraph(); p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runfmt(p.add_run("Panduan eksekusi untuk pembangunan melalui Google Antigravity"), 10, True, NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runfmt(p.add_run("Versi 1.0 • 1 Oktober 2026"), 8.5, color="667085")

# Navigation
doc.add_page_break(); heading(doc, "Status dan cara memakai panduan", 1)
callout(doc, "STATUS: SIAP MENJADI PEDOMAN BUILD", "Arsitektur, alur kerja, struktur data, urutan implementasi, pengujian, deployment, dan operasi sudah ditetapkan. Perhitungan skor final masih harus dikunci dari pemilik proses karena bobot, nilai tiap opsi, pembulatan, dan batas kategori tidak terlihat lengkap pada Form Audit QAS.", LIGHT_AMBER, "C55A11")
body(doc, "Gunakan dokumen ini berurutan. Jangan meminta Antigravity membangun seluruh aplikasi dalam satu prompt. Selesaikan satu tahap, jalankan pemeriksaan, simpan perubahan, lalu lanjut ke tahap berikutnya.")
table(doc, ["Dokumen", "Fungsi", "Cara pakai"], [["Form Audit QAS.pdf", "Sumber struktur pertanyaan dan tampilan form eksisting", "Validasi konten, pilihan jawaban, bukti, dan aturan skor"], ["PRD_Aplikasi_Audit_QAS_Cloudflare.docx", "Kontrak produk dan kebutuhan", "Rujukan fitur, peran, alur, dan acceptance criteria"], ["Panduan ini", "Urutan kerja teknis dan operasional", "Ikuti dari keputusan awal sampai go-live dan operasi"]], [2.25, 2.2, 2.65])
heading(doc, "Definisi tanda", 2)
table(doc, ["Tanda", "Makna"], [["GATE", "Keputusan atau hasil wajib sebelum tahap berikutnya dimulai"], ["OUTPUT", "Artefak yang harus tersedia"], ["CHECK", "Pemeriksaan yang harus lulus"], ["STOP", "Jangan melanjutkan sebelum masalah diselesaikan"]], [1.1, 6.0])
heading(doc, "Daftar isi", 2)
bullets(doc, ["01 Sasaran, ruang lingkup, dan definisi selesai", "02 Keputusan bisnis yang harus dikunci", "03 Arsitektur dan alur data", "04 Persiapan akun dan perangkat kerja", "05 Cara membangun di Antigravity", "06 Membuat fondasi proyek", "07 Menyiapkan Cloudflare D1, R2, dan konfigurasi", "08 Merancang basis data dan data awal", "09 Login, peran, dan pembatasan depo", "10 Implementasi modul aplikasi", "11 Mesin skor dan perbandingan", "12 API, status audit, dan aturan transaksi", "13 Desain antarmuka mobile dan desktop", "14 Offline draft, sinkronisasi, dan bukti foto", "15 Pengujian lengkap", "16 Deployment staging dan production", "17 Pilot, pelatihan, dan rollout tiga depo", "18 Operasi, backup, monitoring, dan insiden", "19 Troubleshooting", "20 Checklist serah terima dan go-live", "21 Paket prompt Antigravity", "22 Referensi resmi"])

# 01
section(doc, 1, "Sasaran, ruang lingkup, dan definisi selesai", "Tetapkan hasil akhir sebelum mulai menulis kode.")
heading(doc, "1.1 Sasaran aplikasi", 2)
bullets(doc, ["PIC QAS Gudang mengisi self audit untuk depo yang menjadi tanggung jawabnya.", "Auditor QAS melakukan audit resmi setelah self audit selesai.", "Jawaban self audit tidak ditampilkan kepada Auditor selama audit resmi masih berjalan agar hasil tetap independen.", "Setelah audit resmi dikirim, sistem membandingkan kedua hasil per pertanyaan, per seksi, per depo, dan per periode.", "Dashboard menampilkan skor, selisih, temuan, progres, bukti, dan status tindak lanjut untuk Karawang, Baros, dan Cirebon.", "Aplikasi dapat digunakan dari browser ponsel dan desktop, dengan draft lokal saat koneksi terganggu."])
heading(doc, "1.2 Ruang lingkup versi pertama", 2)
table(doc, ["Termasuk", "Belum termasuk pada versi pertama"], [["Login, RBAC, tiga depo, master pengguna, template audit berversi", "Integrasi SSO perusahaan di luar Cloudflare Access"], ["Self audit, audit resmi, foto bukti, compare, dashboard, ekspor", "Aplikasi native Android/iOS di toko aplikasi"], ["Draft offline, audit trail, reopen terkontrol, acknowledgement", "Integrasi ERP/WMS bila belum ada API resmi"], ["Deployment Cloudflare dan SOP operasi", "Notifikasi WhatsApp otomatis tanpa penyedia resmi"]], [3.55, 3.55])
heading(doc, "1.3 Definisi selesai", 2)
numbers(doc, ["PIC aktif dapat login, hanya melihat depo yang diizinkan, membuat self audit, menyimpan draft, melampirkan bukti, dan mengirim hasil.", "Auditor dapat memulai audit resmi hanya jika self audit pada siklus yang sama sudah SUBMITTED.", "Auditor tidak dapat membaca jawaban self audit sebelum audit resmi SUBMITTED.", "Perhitungan skor mengikuti aturan resmi yang telah disahkan dan dapat direkonsiliasi dengan contoh manual.", "Dashboard dapat difilter berdasarkan periode, depo, seksi, status, dan jenis audit.", "Semua perubahan kritis tercatat dalam audit trail.", "Aplikasi lulus UAT di Karawang, lalu lulus smoke test di Baros dan Cirebon.", "Backup, restore drill, pengelolaan pengguna, dan prosedur insiden telah diuji."])
callout(doc, "GATE 1", "Pemilik proses menyetujui sasaran, ruang lingkup versi pertama, dan definisi selesai.", LIGHT_GREEN, GREEN)

# 02
section(doc, 2, "Keputusan bisnis yang harus dikunci", "Cegah Antigravity menebak aturan audit yang tidak tertulis.")
heading(doc, "2.1 Keputusan skor — wajib sebelum modul scoring", 2)
table(doc, ["Keputusan", "Yang harus ditetapkan", "Contoh format keputusan"], [["Nilai pilihan", "Nilai untuk setiap jawaban pada tiap pertanyaan", "Baik = …; Cukup = …; Tidak = …"], ["Bobot", "Bobot seksi atau pertanyaan", "J1 = …%; J2 = …%"], ["N/A", "Apakah boleh dan pengaruhnya ke penyebut", "N/A dikeluarkan dari maksimum"], ["Pembulatan", "Jumlah desimal dan metode", "2 desimal; round half-up"], ["Kategori", "Batas skor setiap kategori", "Baik Sekali ≥ …"], ["Finalisasi", "Siapa yang mengesahkan dan dapat mengubah", "Process Owner menyetujui versi skor"]], [1.2, 3.1, 2.8])
callout(doc, "STOP", "Jangan menulis rumus skor produksi dari asumsi. Sistem boleh dibangun dengan konfigurasi placeholder, tetapi tombol publikasi template harus terkunci sampai aturan skor disahkan.", LIGHT_RED, RED)
heading(doc, "2.2 Keputusan proses audit", 2)
table(doc, ["Topik", "Pertanyaan keputusan", "Rekomendasi awal"], [["Siklus", "Bulanan, kuartalan, atau ad hoc?", "Siklus memiliki tanggal mulai, jatuh tempo self audit, dan jatuh tempo audit resmi"], ["Urutan", "Bolehkah Auditor mulai sebelum self audit selesai?", "Tidak; self audit wajib SUBMITTED"], ["Blind audit", "Kapan jawaban self terlihat?", "Setelah audit resmi SUBMITTED"], ["Reopen", "Siapa boleh membuka hasil yang sudah submit?", "Admin/Auditor berwenang dengan alasan wajib"], ["Bukti", "Pertanyaan mana wajib foto/dokumen?", "Atur per pertanyaan di template"], ["Tindak lanjut", "Apakah temuan memerlukan target dan PIC?", "Ya untuk gap di atas ambang yang disepakati"], ["Acknowledgement", "Siapa menyatakan hasil telah dipahami?", "PIC depo setelah hasil resmi diterbitkan"]], [1.2, 3.0, 2.9])
heading(doc, "2.3 Lembar keputusan yang perlu diisi", 2)
table(doc, ["No.", "Keputusan", "Isi final", "Penyetuju"], [[str(i), label, "________________", "________________"] for i, label in enumerate(["Matriks nilai 17 pertanyaan", "Bobot J1 dan J2", "Aturan N/A", "Pembulatan", "Batas kategori", "Bukti wajib", "Jadwal siklus", "Reopen dan approval", "Retensi bukti"], 1)], [.45, 2.55, 2.2, 1.9])
callout(doc, "OUTPUT", "Berita acara keputusan proses dan scoring yang diberi nomor versi. Simpan sebagai referensi template audit.", LIGHT_GREEN, GREEN)

# 03
section(doc, 3, "Arsitektur dan alur data", "Satu deployment Cloudflare untuk antarmuka, API, data, dan bukti.")
heading(doc, "3.1 Komponen", 2)
table(doc, ["Lapisan", "Teknologi", "Tanggung jawab"], [["Antarmuka", "React + TypeScript + Vite PWA", "Tampilan responsif, form audit, dashboard, draft lokal"], ["API", "Cloudflare Worker + Hono", "Validasi, RBAC, aturan proses, perhitungan, audit trail"], ["Data", "Cloudflare D1", "Master, template, siklus, jawaban, skor, event"], ["Berkas", "Cloudflare R2 private bucket", "Foto dan dokumen bukti; tidak dibuka publik"], ["Identitas", "Cloudflare Access", "Login email/IdP; Worker tetap memvalidasi JWT"], ["Offline", "IndexedDB", "Draft jawaban dan antrean unggah sementara"]], [1.1, 2.35, 3.55])
heading(doc, "3.2 Alur permintaan", 2)
code(doc, """Browser mobile / desktop
        │ HTTPS + Access JWT
        ▼
Cloudflare Access ── menolak pengguna tanpa policy
        │
        ▼
Cloudflare Worker + Hono
   ├── verifikasi JWT, role, depot, status audit
   ├── D1: master, audit, jawaban, skor, event
   ├── R2: bukti foto/dokumen private
   └── Static Assets: React PWA
        ▲
        └── IndexedDB: draft dan sync queue""", "ARSITEKTUR LOGIS")
heading(doc, "3.3 Alur bisnis utama", 2)
numbers(doc, ["Admin menyiapkan pengguna, scope depo, template aktif, dan siklus audit.", "PIC memilih siklus, mengisi self audit, melampirkan bukti, lalu submit.", "Sistem mengunci self audit dan mengaktifkan audit resmi untuk Auditor.", "Auditor mengisi audit resmi tanpa melihat jawaban self audit, lalu submit.", "Sistem menghitung skor, membuat snapshot perbandingan, dan membuka hasil compare.", "PIC membaca hasil, menindaklanjuti temuan, lalu memberi acknowledgement.", "Dashboard menggabungkan hasil ketiga depo tanpa mengubah data audit historis."])
callout(doc, "PRINSIP DATA", "Snapshot audit selalu mengacu pada versi template yang dipakai saat audit dibuat. Perubahan template baru tidak boleh mengubah audit lama.", CREAM, NAVY)

# 04
section(doc, 4, "Persiapan akun dan perangkat kerja", "Siapkan akses sebelum membuat resource cloud.")
heading(doc, "4.1 Akun", 2)
table(doc, ["Kebutuhan", "Keterangan", "Status"], [["Google Antigravity", "Lingkungan utama untuk membangun dan meninjau perubahan", "☐"], ["Cloudflare", "Workers, D1, R2, Access, DNS", "☐"], ["Domain aktif di Cloudflare", "Diperlukan untuk hostname Access yang stabil; gunakan domain perusahaan yang sudah ada", "☐"], ["Repository Git", "Disarankan untuk riwayat versi dan rollback", "☐"], ["Daftar email pengguna", "Email, nama, peran, dan depo", "☐"]], [2.0, 4.6, .5])
body(doc, "Catatan biaya: target dapat berjalan pada free tier untuk pilot dengan volume kecil. Domain, kapasitas aktual, dan penggunaan di atas kuota tetap dapat menimbulkan biaya. Pantau konsumsi dari dashboard Cloudflare.")
heading(doc, "4.2 Perangkat pengembangan", 2)
bullets(doc, ["Node.js versi LTS dan npm.", "Git untuk version control.", "Chrome atau Edge untuk pengujian desktop dan mobile emulation.", "Ponsel Android/iPhone nyata untuk pengujian kamera, upload, offline, dan instalasi PWA.", "Akses ke Form Audit QAS dan PRD final."])
heading(doc, "4.3 Konvensi nama", 2)
table(doc, ["Objek", "Staging", "Production"], [["Worker", "qas-audit-staging", "qas-audit-prod"], ["D1", "qas-audit-staging-db", "qas-audit-prod-db"], ["R2", "qas-audit-staging-evidence", "qas-audit-prod-evidence"], ["Hostname", "qas-staging.<domain>", "qas.<domain>"]], [2.0, 2.55, 2.55])
callout(doc, "GATE 2", "Cloudflare, domain, Antigravity, daftar pengguna, dan hak admin tersedia. Nama resource disetujui.", LIGHT_GREEN, GREEN)

# 05
section(doc, 5, "Cara membangun di Antigravity", "Gunakan pekerjaan kecil dengan acceptance criteria yang dapat diuji.")
heading(doc, "5.1 Membuka workspace", 2)
numbers(doc, ["Buat folder kerja qas-audit-app dan buka folder tersebut di Antigravity.", "Salin PRD dan panduan ini ke docs/. Jangan masukkan PDF sensitif ke repository publik.", "Buat repository Git dan commit kondisi awal.", "Tambahkan AGENTS.md agar Antigravity selalu mengikuti batasan proyek.", "Kerjakan milestone sesuai urutan pada Bab 21. Satu prompt menghasilkan satu perubahan yang dapat ditinjau."])
heading(doc, "5.2 Isi AGENTS.md yang disarankan", 2)
code(doc, """# Aplikasi Audit QAS

- Gunakan React + TypeScript + Vite dan Cloudflare Worker + Hono.
- Gunakan D1 untuk data relasional dan R2 private untuk bukti.
- Semua input divalidasi di server dengan Zod.
- Jangan percaya role, depot, skor, atau status yang dikirim browser.
- Auditor tidak boleh melihat jawaban self audit sebelum official submitted.
- Semua perubahan status kritis harus masuk audit_events.
- Template dan scoring harus versioned; audit lama immutable.
- Jangan hardcode aturan skor sebelum keputusan bisnis disahkan.
- Setiap perubahan harus lulus lint, typecheck, dan tes relevan.
- Gunakan Bahasa Indonesia untuk UI pengguna.
- Jangan commit .dev.vars, token, secret, atau data audit nyata.""", "AGENTS.md")
heading(doc, "5.3 Siklus kerja setiap milestone", 2)
table(doc, ["Langkah", "Yang dilakukan", "Bukti selesai"], [["1. Minta", "Kirim prompt dengan ruang lingkup jelas", "Daftar file yang akan berubah"], ["2. Tinjau", "Periksa rencana sebelum perubahan besar", "Tidak ada scope creep"], ["3. Jalankan", "Izinkan perubahan dan dependency relevan", "Aplikasi dapat dijalankan"], ["4. Uji", "Lint, typecheck, unit/integration/E2E relevan", "Semua pemeriksaan lulus"], ["5. Inspeksi", "Lihat UI mobile dan desktop", "Tidak ada teks terpotong/alur buntu"], ["6. Simpan", "Commit milestone", "Riwayat dapat di-rollback"]], [.9, 3.6, 2.6])
callout(doc, "STOP", "Jika Antigravity mengganti arsitektur, menghapus tes, membuka bucket R2, atau menebak rumus skor, hentikan dan arahkan kembali ke PRD serta AGENTS.md.", LIGHT_RED, RED)

# 06
section(doc, 6, "Membuat fondasi proyek", "Scaffold React dan Worker dalam satu proyek.")
heading(doc, "6.1 Membuat project", 2)
body(doc, "Jalankan perintah berikut dari terminal Antigravity. Saat wizard meminta pilihan, gunakan TypeScript, Git, dan deployment Workers.")
code(doc, """npm create cloudflare@latest -- qas-audit-app --framework=react
cd qas-audit-app
npm install
npm install hono zod jose idb react-router-dom
npm install -D vitest @testing-library/react @testing-library/jest-dom @playwright/test eslint prettier""", "POWERSHELL / TERMINAL")
heading(doc, "6.2 Struktur folder target", 2)
code(doc, """qas-audit-app/
├─ docs/                     PRD, keputusan scoring, panduan
├─ migrations/               migrasi D1 berurutan
├─ scripts/                  seed dan pemeriksaan data
├─ src/
│  ├─ app/                   router, provider, layout
│  ├─ components/            komponen UI umum
│  ├─ features/              auth, audit, compare, dashboard, master
│  ├─ lib/                   API client, schema, formatter, offline
│  └─ styles/
├─ worker/
│  ├─ index.ts               entry Worker
│  ├─ middleware/            auth, RBAC, error, request id
│  ├─ routes/                endpoint API
│  ├─ services/              aturan bisnis
│  └─ repositories/          akses D1/R2
├─ tests/                    unit, integration, E2E
├─ public/                   manifest dan icon PWA
├─ AGENTS.md
├─ wrangler.jsonc
└─ package.json""", "STRUKTUR DIREKTORI")
heading(doc, "6.3 Script dan pemeriksaan awal", 2)
code(doc, """npm run dev
npm run build
npm run lint
npm run typecheck
npm run test""", "CHECK")
bullets(doc, ["Halaman starter tampil tanpa error di desktop.", "Emulasi lebar 360 px tidak memunculkan scroll horizontal.", "Build produksi berhasil.", "Tidak ada token atau file .dev.vars masuk Git."])
callout(doc, "GATE 3", "Scaffold berjalan lokal, build bersih, dan struktur folder disetujui.", LIGHT_GREEN, GREEN)

# 07
section(doc, 7, "Menyiapkan Cloudflare D1, R2, dan konfigurasi", "Pisahkan data terstruktur dari bukti foto dan dokumen.")
heading(doc, "7.1 Login CLI", 2); code(doc, "npx wrangler login\nnpx wrangler whoami", "TERMINAL")
heading(doc, "7.2 Membuat resource", 2)
code(doc, """npx wrangler d1 create qas-audit-staging-db --location apac
npx wrangler d1 create qas-audit-prod-db --location apac
npx wrangler r2 bucket create qas-audit-staging-evidence
npx wrangler r2 bucket create qas-audit-prod-evidence""", "TERMINAL")
body(doc, "Simpan database_id yang ditampilkan oleh Wrangler. Jangan menyalin ID staging ke konfigurasi production.")
heading(doc, "7.3 Contoh wrangler.jsonc", 2)
code(doc, """{
  "$schema": "node_modules/wrangler/config-schema.json",
  "name": "qas-audit-prod",
  "main": "worker/index.ts",
  "compatibility_date": "2026-10-01",
  "assets": {
    "not_found_handling": "single-page-application",
    "run_worker_first": ["/api/*"]
  },
  "d1_databases": [{
    "binding": "DB",
    "database_name": "qas-audit-prod-db",
    "database_id": "GANTI_DENGAN_ID_D1_PRODUCTION",
    "migrations_dir": "migrations"
  }],
  "r2_buckets": [{
    "binding": "EVIDENCE",
    "bucket_name": "qas-audit-prod-evidence"
  }],
  "vars": { "APP_ENV": "production", "MAX_EVIDENCE_BYTES": "5242880" }
}""", "wrangler.jsonc")
body(doc, "Untuk staging, gunakan environment atau file konfigurasi terpisah. Pada Cloudflare Vite plugin, pilih environment saat dev/build melalui CLOUDFLARE_ENV sesuai dokumentasi versi yang digunakan.")
heading(doc, "7.4 Aturan penyimpanan bukti", 2)
bullets(doc, ["Bucket R2 private; jangan aktifkan public development URL untuk bukti audit.", "Browser mengunggah ke Worker; Worker memeriksa role, audit, MIME, ukuran, dan nama objek.", "Object key menggunakan ID acak: evidence/{depotId}/{auditId}/{answerId}/{uuid}.jpg.", "Metadata D1 menyimpan nama asli, MIME, byte, hash, uploader, dan waktu upload.", "Download bukti melewati endpoint yang mengulang pemeriksaan akses."])
callout(doc, "CHECK", "Resource staging dan production terlihat di dashboard. Binding DB dan EVIDENCE sama di kode, tetapi menunjuk resource berbeda.", LIGHT_GREEN, GREEN)

# 08
section(doc, 8, "Merancang basis data dan data awal", "Simpan versi template dan audit historis secara utuh.")
heading(doc, "8.1 Entitas inti", 2)
table(doc, ["Entitas", "Fungsi", "Aturan utama"], [
    ["depots", "Karawang, Baros, Cirebon", "Kode unik; nonaktif tidak menghapus riwayat"],
    ["users / user_role_scopes", "Email, role, scope depo", "Satu email dapat memiliki lebih dari satu scope"],
    ["audit_templates / versions", "Form dan versi scoring", "Versi published immutable"],
    ["sections / questions / options", "J1, J2, 17 pertanyaan, pilihan", "Urutan eksplisit; bukti wajib per pertanyaan"],
    ["audit_cycles", "Periode dan tenggat", "Unik per nama/periode"],
    ["audits / audit_answers", "Self atau official dan jawabannya", "Satu audit per tipe, depo, siklus"],
    ["evidence_files", "Metadata objek R2", "Kunci objek private dan hash"],
    ["comparison_snapshots", "Hasil compare saat official submit", "Tidak dihitung ulang diam-diam"],
    ["audit_events", "Audit trail", "Append-only"],
], [1.8, 2.55, 2.75])

heading(doc, "8.2 Migrasi awal — master dan template", 2)
code(doc, """PRAGMA foreign_keys = ON;

CREATE TABLE depots (
  id TEXT PRIMARY KEY,
  code TEXT NOT NULL UNIQUE,
  name TEXT NOT NULL,
  is_active INTEGER NOT NULL DEFAULT 1 CHECK (is_active IN (0,1)),
  created_at TEXT NOT NULL
);

CREATE TABLE users (
  id TEXT PRIMARY KEY,
  email TEXT NOT NULL UNIQUE COLLATE NOCASE,
  full_name TEXT NOT NULL,
  is_active INTEGER NOT NULL DEFAULT 1 CHECK (is_active IN (0,1)),
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL
);

CREATE TABLE user_role_scopes (
  id TEXT PRIMARY KEY,
  user_id TEXT NOT NULL REFERENCES users(id),
  role TEXT NOT NULL CHECK (role IN ('ADMIN','PIC_QAS','AUDITOR_QAS','VIEWER')),
  depot_id TEXT REFERENCES depots(id),
  UNIQUE(user_id, role, depot_id)
);

CREATE TABLE audit_templates (
  id TEXT PRIMARY KEY,
  code TEXT NOT NULL UNIQUE,
  name TEXT NOT NULL,
  is_active INTEGER NOT NULL DEFAULT 1
);

CREATE TABLE audit_template_versions (
  id TEXT PRIMARY KEY,
  template_id TEXT NOT NULL REFERENCES audit_templates(id),
  version_no INTEGER NOT NULL,
  status TEXT NOT NULL CHECK (status IN ('DRAFT','PUBLISHED','RETIRED')),
  scoring_config_json TEXT,
  published_at TEXT,
  published_by TEXT REFERENCES users(id),
  UNIQUE(template_id, version_no)
);""", "migrations/0001_initial.sql — BAGIAN 1")

heading(doc, "8.3 Migrasi awal — pertanyaan dan siklus", 2)
code(doc, """CREATE TABLE audit_sections (
  id TEXT PRIMARY KEY,
  version_id TEXT NOT NULL REFERENCES audit_template_versions(id),
  code TEXT NOT NULL,
  title TEXT NOT NULL,
  display_order INTEGER NOT NULL,
  weight REAL,
  UNIQUE(version_id, code)
);

CREATE TABLE audit_questions (
  id TEXT PRIMARY KEY,
  section_id TEXT NOT NULL REFERENCES audit_sections(id),
  code TEXT NOT NULL,
  prompt TEXT NOT NULL,
  display_order INTEGER NOT NULL,
  evidence_required INTEGER NOT NULL DEFAULT 0,
  is_required INTEGER NOT NULL DEFAULT 1,
  weight REAL,
  UNIQUE(section_id, code)
);

CREATE TABLE answer_options (
  id TEXT PRIMARY KEY,
  question_id TEXT NOT NULL REFERENCES audit_questions(id),
  code TEXT NOT NULL,
  label TEXT NOT NULL,
  numeric_value REAL,
  display_order INTEGER NOT NULL,
  is_na INTEGER NOT NULL DEFAULT 0,
  UNIQUE(question_id, code)
);

CREATE TABLE audit_cycles (
  id TEXT PRIMARY KEY,
  code TEXT NOT NULL UNIQUE,
  title TEXT NOT NULL,
  period_start TEXT NOT NULL,
  period_end TEXT NOT NULL,
  self_due_at TEXT NOT NULL,
  official_due_at TEXT NOT NULL,
  status TEXT NOT NULL CHECK (status IN ('DRAFT','OPEN','CLOSED')),
  template_version_id TEXT NOT NULL REFERENCES audit_template_versions(id)
);""", "migrations/0001_initial.sql — BAGIAN 2")

heading(doc, "8.4 Migrasi awal — transaksi", 2)
code(doc, """CREATE TABLE audits (
  id TEXT PRIMARY KEY,
  cycle_id TEXT NOT NULL REFERENCES audit_cycles(id),
  depot_id TEXT NOT NULL REFERENCES depots(id),
  audit_type TEXT NOT NULL CHECK (audit_type IN ('SELF','OFFICIAL')),
  status TEXT NOT NULL CHECK (status IN ('DRAFT','SUBMITTED','REOPENED','VOID')),
  assigned_user_id TEXT NOT NULL REFERENCES users(id),
  template_version_id TEXT NOT NULL REFERENCES audit_template_versions(id),
  score REAL,
  category TEXT,
  submitted_at TEXT,
  created_at TEXT NOT NULL,
  updated_at TEXT NOT NULL,
  UNIQUE(cycle_id, depot_id, audit_type)
);

CREATE TABLE audit_answers (
  id TEXT PRIMARY KEY,
  audit_id TEXT NOT NULL REFERENCES audits(id),
  question_id TEXT NOT NULL REFERENCES audit_questions(id),
  option_id TEXT REFERENCES answer_options(id),
  note TEXT,
  numeric_value_snapshot REAL,
  updated_at TEXT NOT NULL,
  UNIQUE(audit_id, question_id)
);

CREATE TABLE evidence_files (
  id TEXT PRIMARY KEY,
  answer_id TEXT NOT NULL REFERENCES audit_answers(id),
  object_key TEXT NOT NULL UNIQUE,
  original_name TEXT NOT NULL,
  mime_type TEXT NOT NULL,
  size_bytes INTEGER NOT NULL,
  sha256 TEXT,
  uploaded_by TEXT NOT NULL REFERENCES users(id),
  uploaded_at TEXT NOT NULL,
  deleted_at TEXT
);

CREATE TABLE comparison_snapshots (
  id TEXT PRIMARY KEY,
  cycle_id TEXT NOT NULL REFERENCES audit_cycles(id),
  depot_id TEXT NOT NULL REFERENCES depots(id),
  self_audit_id TEXT NOT NULL REFERENCES audits(id),
  official_audit_id TEXT NOT NULL REFERENCES audits(id),
  summary_json TEXT NOT NULL,
  created_at TEXT NOT NULL,
  UNIQUE(cycle_id, depot_id)
);""", "migrations/0001_initial.sql — BAGIAN 3")

heading(doc, "8.5 Migrasi awal — audit trail dan indeks", 2)
code(doc, """CREATE TABLE audit_events (
  id TEXT PRIMARY KEY,
  entity_type TEXT NOT NULL,
  entity_id TEXT NOT NULL,
  event_type TEXT NOT NULL,
  actor_user_id TEXT REFERENCES users(id),
  reason TEXT,
  payload_json TEXT,
  request_id TEXT,
  created_at TEXT NOT NULL
);

CREATE TABLE acknowledgements (
  id TEXT PRIMARY KEY,
  official_audit_id TEXT NOT NULL REFERENCES audits(id),
  user_id TEXT NOT NULL REFERENCES users(id),
  note TEXT,
  acknowledged_at TEXT NOT NULL,
  UNIQUE(official_audit_id, user_id)
);

CREATE INDEX idx_audits_cycle_depot ON audits(cycle_id, depot_id);
CREATE INDEX idx_answers_audit ON audit_answers(audit_id);
CREATE INDEX idx_events_entity ON audit_events(entity_type, entity_id, created_at);
CREATE INDEX idx_evidence_answer ON evidence_files(answer_id, deleted_at);""", "migrations/0001_initial.sql — BAGIAN 4")

heading(doc, "8.6 Membuat dan menjalankan migrasi", 2)
code(doc, """npx wrangler d1 migrations create DB initial_schema
npx wrangler d1 migrations apply DB --local
npm run dev

# Setelah tes lokal lulus:
npx wrangler d1 migrations list DB --remote
npx wrangler d1 migrations apply DB --remote""", "TERMINAL")
heading(doc, "8.7 Seed minimum", 2)
bullets(doc, ["Depo: KRW/Karawang, BRS/Baros, CRB/Cirebon.", "Pengguna pertama: satu Admin, minimal satu PIC per depo, minimal satu Auditor.", "Template QAS versi 1 dalam status DRAFT.", "Seksi J1 dan J2 serta 17 pertanyaan dari formulir sumber.", "Pilihan jawaban dan scoring_config_json tetap kosong atau belum disahkan sampai keputusan scoring tersedia."])
callout(doc, "CHECK", "Hitung: 3 depo, 17 pertanyaan, tidak ada code duplikat, dan tidak ada versi PUBLISHED tanpa konfigurasi skor lengkap.", LIGHT_GREEN, GREEN)

# 09
section(doc, 9, "Login, peran, dan pembatasan depo", "Cloudflare Access mengidentifikasi pengguna; aplikasi menentukan kewenangannya.", False)
heading(doc, "9.1 Menyiapkan Cloudflare Access", 2)
numbers(doc, ["Buka Cloudflare Zero Trust dan buat team name.", "Aktifkan One-time PIN atau identity provider perusahaan.", "Deploy staging sementara agar hostname tersedia.", "Buat Access application Self-hosted and private, lalu tambahkan public hostname qas-staging.<domain>.", "Buat Allow policy hanya untuk email atau domain organisasi yang disetujui.", "Salin Team Domain dan Application Audience (AUD) Tag.", "Simpan nilai sensitif sebagai Worker secret, bukan source code."])
code(doc, "npx wrangler secret put ACCESS_TEAM_DOMAIN\nnpx wrangler secret put ACCESS_AUD", "TERMINAL")
heading(doc, "9.2 Verifikasi JWT di Worker", 2)
bullets(doc, ["Ambil token dari header Cf-Access-Jwt-Assertion.", "Validasi signature menggunakan JWKS, issuer, audience, dan expiry.", "Ambil email tervalidasi dari claim token.", "Cari email di users dan user_role_scopes.", "Tolak 401 jika token tidak valid; 403 jika pengguna tidak aktif atau scope tidak sesuai.", "Tambahkan request ID dan catat aksi sensitif ke audit_events."])
callout(doc, "ATURAN KEAMANAN", "Header email tanpa validasi signature tidak boleh dipercaya. Access melindungi pintu masuk, sedangkan Worker tetap memvalidasi JWT dan aturan peran/depo pada setiap endpoint.", LIGHT_RED, RED)
heading(doc, "9.3 Matriks hak akses", 2)
table(doc, ["Aksi", "Admin", "PIC QAS", "Auditor QAS", "Viewer"], [["Kelola pengguna/depo", "Ya", "Tidak", "Tidak", "Tidak"], ["Kelola draft template", "Ya", "Tidak", "Opsional", "Tidak"], ["Buat/isi self audit", "Tidak", "Depo sendiri", "Tidak", "Tidak"], ["Buat/isi official audit", "Tidak", "Tidak", "Scope audit", "Tidak"], ["Lihat self sebelum official submit", "Terbatas", "Milik depo", "Tidak", "Tidak"], ["Lihat hasil compare", "Ya", "Depo sendiri", "Ya", "Sesuai scope"], ["Reopen", "Dengan alasan", "Tidak", "Jika diberi hak", "Tidak"]], [2.45, .9, 1.25, 1.3, 1.2], 7.8)
heading(doc, "9.4 Tes otorisasi minimum", 2)
bullets(doc, ["PIC Karawang tidak dapat membuka audit Baros melalui URL atau request.", "Auditor tidak dapat membaca endpoint self answers sebelum official submit.", "Pengguna yang dinonaktifkan di D1 ditolak walaupun login Access berhasil.", "Viewer tidak dapat POST, PATCH, DELETE, submit, atau upload.", "Setiap 403 memiliki request ID tanpa membocorkan data audit."])
callout(doc, "GATE 4", "Login staging berhasil dan seluruh tes negatif RBAC lulus sebelum modul audit dibangun.", LIGHT_GREEN, GREEN)

# 10
section(doc, 10, "Implementasi modul aplikasi", "Bangun dari master dan template menuju audit, compare, lalu dashboard.", False)
modules = [
    ("10.1 Kerangka aplikasi dan profil", ["Layout desktop: sidebar, header konteks depo/siklus, area konten.", "Layout mobile: top bar ringkas, bottom navigation, target sentuh minimal 44 px.", "Halaman /me menampilkan nama, email, role, dan scope depo dari server.", "Error boundary, loading skeleton, empty state, toast, dan halaman 403/404."], "Login mengarah ke dashboard sesuai role; tidak ada data dummy tersisa."),
    ("10.2 Master pengguna dan scope", ["Admin mengatur nama, email, status aktif, role, dan depo.", "Email disimpan lowercase dan unik.", "Menonaktifkan pengguna tidak menghapus audit historis.", "Semua perubahan master tercatat di audit_events."], "PIC tiga depo dan Auditor dapat diuji dengan akun terpisah."),
    ("10.3 Template audit berversi", ["Admin membuat draft versi, seksi, pertanyaan, pilihan, urutan, bukti wajib, dan bobot.", "Validasi menolak publish bila scoring belum lengkap atau code duplikat.", "Versi PUBLISHED tidak dapat diedit; perubahan melalui clone ke versi baru.", "Preview menunjukkan urutan yang dilihat pengguna."], "Template sumber memiliki J1, J2, dan 17 pertanyaan terverifikasi."),
    ("10.4 Siklus audit", ["Admin membuat judul/periode, due date self/official, dan template version.", "Saat OPEN, sistem membuat slot self dan official per depo.", "Siklus CLOSED read-only kecuali reopen terotorisasi."], "Satu siklus uji memiliki enam audit untuk tiga depo."),
    ("10.5 Self audit", ["Form dikelompokkan per seksi dan menampilkan progres answered/required.", "Autosave lokal cepat; simpan server berkala dan saat pindah langkah.", "Submit memeriksa pertanyaan wajib, bukti wajib, dan konflik versi.", "Setelah submit, jawaban terkunci dan official audit tersedia."], "PIC dapat menutup browser, membuka kembali, melanjutkan draft, lalu submit."),
    ("10.6 Audit resmi", ["Auditor melihat status self selesai tanpa jawaban, nilai, catatan, atau bukti self.", "Form resmi memakai template version yang sama.", "Submit transaksional: validasi, snapshot nilai, skor, compare, dan event."], "Tes API membuktikan blind audit sampai official SUBMITTED."),
    ("10.7 Hasil compare dan acknowledgement", ["Tampilkan self vs official per pertanyaan, selisih, catatan, dan bukti sesuai izin.", "Ringkas jumlah sama, berbeda, lebih tinggi, lebih rendah, dan gap terbesar.", "PIC memberi acknowledgement dengan waktu, pengguna, dan catatan opsional."], "Snapshot compare tidak berubah bila template baru diterbitkan."),
    ("10.8 Dashboard dan ekspor", ["Kartu status; tren skor; perbandingan depo; gap per seksi; temuan prioritas.", "Filter periode, depo, seksi, status, dan jenis audit diterapkan server-side.", "Ekspor CSV memakai filter sama dan mencantumkan waktu ekspor."], "Kartu, chart, daftar, dan CSV konsisten untuk filter yang sama."),
]
for title, items, done in modules:
    heading(doc, title, 2); bullets(doc, items); callout(doc, "SELESAI JIKA", done, LIGHT_GREEN, GREEN)

# 11
section(doc, 11, "Mesin skor dan perbandingan", "Konfigurasi skor menjadi data versi, bukan angka tersembunyi di kode.")
heading(doc, "11.1 Struktur konfigurasi", 2)
code(doc, """{
  "method": "weighted_normalized",
  "decimalPlaces": 2,
  "rounding": "half_up",
  "naPolicy": "exclude_from_denominator",
  "categories": [
    { "code": "TBD", "label": "Menunggu keputusan", "min": null }
  ]
}""", "CONTOH scoring_config_json — PLACEHOLDER")
body(doc, "Struktur tersebut bukan aturan skor final. Setelah keputusan bisnis disahkan, masukkan pilihan, bobot, N/A, pembulatan, dan kategori yang benar ke versi template.")
heading(doc, "11.2 Urutan perhitungan yang aman", 2)
numbers(doc, ["Ambil answer option dan numeric_value dari versi template audit.", "Validasi semua jawaban wajib dan kebijakan N/A.", "Hitung skor mentah per pertanyaan dan seksi menggunakan bobot versi.", "Normalisasi hanya bila metode resmi meminta.", "Terapkan pembulatan satu kali pada tahap yang ditetapkan.", "Tentukan kategori dari tabel ambang versi.", "Simpan nilai per jawaban/seksi, skor final, kategori, dan configuration hash sebagai snapshot."])
heading(doc, "11.3 Tes golden case", 2)
table(doc, ["Kasus", "Input", "Expected"], [["Semua maksimum", "Semua opsi nilai tertinggi", "Skor maksimum dan kategori tertinggi"], ["Semua minimum", "Semua opsi nilai terendah", "Skor minimum dan kategori terendah"], ["N/A", "Satu pertanyaan N/A", "Penyebut sesuai keputusan"], ["Batas kategori", "Nilai tepat di ambang", "Kategori sesuai inclusive/exclusive"], ["Pembulatan", "Nilai pecahan kritis", "Sama dengan contoh manual"], ["Rekonsiliasi", "Contoh formulir 4,37", "Cocok setelah aturan resmi diketahui"]], [1.2, 3.1, 2.8])
callout(doc, "GATE 5", "Process Owner menandatangani hasil golden case. Tanpa persetujuan ini, skor produksi dan kategori tidak boleh diaktifkan.", LIGHT_AMBER, "C55A11")

# 12
section(doc, 12, "API, status audit, dan aturan transaksi", "Validasi bisnis berada di server agar tidak dapat dilewati dari browser.")
heading(doc, "12.1 Endpoint minimum", 2)
table(doc, ["Metode", "Endpoint", "Fungsi / akses"], [["GET", "/api/me", "Profil dan scope"], ["GET", "/api/cycles", "Siklus sesuai scope"], ["POST", "/api/admin/cycles", "Buat siklus — Admin"], ["GET", "/api/audits/:id", "Metadata; field difilter role/status"], ["GET", "/api/audits/:id/form", "Template dan jawaban yang diizinkan"], ["PATCH", "/api/audits/:id/answers", "Upsert draft + version token"], ["POST", "/api/audits/:id/evidence", "Upload bukti private"], ["GET", "/api/evidence/:id", "Streaming setelah otorisasi"], ["POST", "/api/audits/:id/submit", "Submit idempotent"], ["POST", "/api/audits/:id/reopen", "Reopen dengan alasan"], ["GET", "/api/comparisons/:cycleId/:depotId", "Compare setelah official submit"], ["GET", "/api/dashboard", "Agregasi filter/scope"], ["GET", "/api/exports/audits.csv", "Ekspor filter/scope"]], [.65, 2.8, 3.65], 7.8)
heading(doc, "12.2 State machine", 2)
code(doc, """SELF:     DRAFT ──submit──> SUBMITTED ──reopen──> REOPENED
                                      │                 │
                                      └─ official open  └─submit ulang

OFFICIAL: DRAFT (terkunci sampai self submitted)
          └─ mulai/isi ──submit──> SUBMITTED ──reopen──> REOPENED

SIKLUS:   DRAFT ──open──> OPEN ──close──> CLOSED""", "TRANSISI YANG DIIZINKAN")
heading(doc, "12.3 Aturan transaksi submit", 2)
bullets(doc, ["Gunakan idempotency key agar klik ganda tidak membuat event/snapshot ganda.", "Periksa updated_at/version token agar draft lama tidak menimpa draft baru.", "Validasi status, role, depot, template, jawaban wajib, dan bukti.", "Official submit menghitung skor dan compare dalam satu unit logis.", "Jika satu langkah gagal, status tidak berubah menjadi SUBMITTED.", "Gunakan kode error stabil: AUTH_REQUIRED, FORBIDDEN, INVALID_STATE, VALIDATION_FAILED, CONFLICT, RATE_LIMITED, INTERNAL_ERROR."])
heading(doc, "12.4 Bentuk respons", 2)
code(doc, """// sukses
{ "data": { ... }, "requestId": "req_..." }

// gagal
{
  "error": {
    "code": "VALIDATION_FAILED",
    "message": "Periksa data yang wajib diisi.",
    "fields": { "question.J1-03": "Bukti wajib dilampirkan" }
  },
  "requestId": "req_..."
}""", "JSON")

# 13
section(doc, 13, "Desain antarmuka mobile dan desktop", "Prioritaskan pengisian cepat, status jelas, dan sedikit langkah.", False)
heading(doc, "13.1 Peta halaman", 2)
table(doc, ["Rute", "Pengguna", "Isi utama"], [["/", "Semua", "Dashboard sesuai role"], ["/cycles", "Semua", "Periode dan progres"], ["/audits/:id", "PIC/Auditor", "Form per seksi"], ["/compare/:cycle/:depot", "Sesuai scope", "Self vs official"], ["/findings", "Sesuai scope", "Temuan/tindak lanjut"], ["/admin/users", "Admin", "Pengguna dan scope"], ["/admin/templates", "Admin", "Template dan versi"], ["/admin/cycles", "Admin", "Siklus"]], [2.1, 1.6, 3.5])
heading(doc, "13.2 Pola form audit mobile", 2)
bullets(doc, ["Satu kartu pertanyaan pada satu waktu atau daftar pendek per seksi; hindari tabel lebar.", "Judul tetap terlihat, pilihan besar, tombol bukti dekat jawaban.", "Progress bar menunjukkan answered/required dan bukti kurang.", "Footer lengket: Sebelumnya, Simpan, Berikutnya; Submit setelah review.", "Kamera/galeri dengan preview dan ukuran file.", "Status offline, pending, syncing, dan tersimpan terlihat jelas."])
heading(doc, "13.3 Pola desktop", 2)
bullets(doc, ["Sidebar modul dan breadcrumb siklus/depo/audit.", "Daftar pertanyaan di kiri, form tengah, progres/bukti kanan bila ruang cukup.", "Dashboard memakai grid kartu, chart, dan tabel detail.", "Filter di atas dan dapat dibagikan via query string tanpa data sensitif."])
heading(doc, "13.4 Aturan kualitas UI", 2)
table(doc, ["Area", "Aturan"], [["Teks", "Bahasa Indonesia, label singkat, error menjelaskan cara memperbaiki"], ["Warna", "Jangan mengandalkan warna saja; sertakan ikon/label"], ["Ukuran", "Body nyaman dibaca; target sentuh minimal 44 px"], ["Form", "Label terlihat; fokus keyboard jelas; field wajib ditandai"], ["Tabel", "Mobile menjadi card/list; tanpa scroll horizontal membingungkan"], ["Chart", "Nilai tersedia di tabel/tooltip; sumbu dan unit jelas"]], [1.2, 5.9])
callout(doc, "CHECK", "Uji minimal pada 360×800, 390×844, 768×1024, 1366×768, dan 1920×1080 serta satu ponsel nyata.", LIGHT_GREEN, GREEN)

# 14
section(doc, 14, "Offline draft, sinkronisasi, dan bukti foto", "Draft tidak hilang saat koneksi terputus, tetapi server tetap sumber kebenaran.", False)
heading(doc, "14.1 Data lokal", 2)
table(doc, ["Store IndexedDB", "Isi", "Kebijakan"], [["auditDrafts", "Jawaban, catatan, localUpdatedAt, serverVersion", "Hapus setelah submit"], ["uploadQueue", "Blob, MIME, ukuran, answerId, retry count", "Batasi total byte dan tampilkan antrean"], ["syncQueue", "Operasi upsert jawaban", "Idempotency key per operasi"], ["referenceCache", "Template aktif dan master ringan", "Versi dan expiry wajib"]], [1.45, 3.1, 2.55])
heading(doc, "14.2 Algoritma sinkronisasi", 2)
numbers(doc, ["Perubahan ditulis ke IndexedDB dahulu dengan status pending.", "Jika online, kirim batch jawaban dengan serverVersion terakhir.", "Server menerima bila versi cocok lalu mengembalikan versi baru.", "Jika berbeda, tandai conflict dan minta pilihan; jangan timpa otomatis.", "Unggah bukti setelah jawaban punya ID server; gunakan retry bertahap.", "Submit hanya aktif bila semua operasi dan bukti wajib tersinkron."])
heading(doc, "14.3 Pengolahan foto", 2)
bullets(doc, ["Terima JPEG/PNG dan format lain hanya bila perlu.", "Kompresi di browser tanpa menghilangkan keterbacaan objek.", "Server memeriksa MIME berdasarkan isi/header, bukan ekstensi saja.", "Tolak file di atas MAX_EVIDENCE_BYTES dengan pesan jelas.", "Simpan SHA-256 untuk deteksi duplikat dan integritas dasar.", "Bersihkan objek orphan secara aman setelah masa tunggu."])
heading(doc, "14.4 Skenario offline wajib", 2)
bullets(doc, ["Buka online, matikan jaringan, isi 5 jawaban, tutup tab, buka kembali; draft tetap ada.", "Ambil foto offline; saat online foto tersinkron.", "Edit di dua perangkat; sistem mendeteksi konflik.", "Token Access kedaluwarsa; login ulang tanpa menghapus draft.", "Koneksi putus saat submit; retry tidak membuat submit ganda."])

# 15
section(doc, 15, "Pengujian lengkap", "Uji aturan bisnis, keamanan, integrasi cloud, dan pengalaman pengguna.")
heading(doc, "15.1 Piramida tes", 2)
table(doc, ["Lapisan", "Fokus", "Contoh"], [["Unit", "Fungsi murni", "scoring, kategori, state, validation"], ["Integration", "Worker + D1/R2 lokal", "RBAC, submit, upload, compare, trail"], ["E2E", "Browser terhadap staging", "login, self, official, compare"], ["UAT", "Proses nyata", "PIC dan Auditor satu siklus"], ["Security", "Penyalahgunaan", "IDOR, tampering, invalid upload, blind leak"]], [1.1, 2.1, 3.9])
heading(doc, "15.2 Perintah pemeriksaan", 2); code(doc, "npm run lint\nnpm run typecheck\nnpm run test\nnpm run build\nnpm run test:e2e", "QUALITY GATE")
heading(doc, "15.3 Matriks UAT inti", 2)
table(doc, ["ID", "Skenario", "Expected"], [["UAT-01", "PIC Karawang login", "Hanya Karawang terlihat"], ["UAT-02", "Simpan draft lalu logout", "Draft dapat dilanjutkan"], ["UAT-03", "Submit tanpa bukti wajib", "Ditolak dan lokasi error jelas"], ["UAT-04", "Self submit lengkap", "SUBMITTED dan terkunci"], ["UAT-05", "Auditor buka official", "Self answers tidak terlihat"], ["UAT-06", "Official submit", "Skor dan compare dibuat sekali"], ["UAT-07", "PIC lihat compare", "Detail/ringkasan konsisten"], ["UAT-08", "Viewer PATCH", "403"], ["UAT-09", "PIC ubah depotId", "403 tanpa kebocoran"], ["UAT-10", "Admin reopen dengan alasan", "Status/event benar"], ["UAT-11", "Offline lalu sync", "Tidak hilang atau ganda"], ["UAT-12", "Ekspor Cirebon", "CSV hanya Cirebon; angka cocok"]], [.7, 3.0, 3.4], 7.7)
heading(doc, "15.4 Data uji", 2)
bullets(doc, ["Akun uji: admin, pic.karawang, pic.baros, pic.cirebon, auditor, viewer.", "Dua siklus: satu OPEN dan satu CLOSED.", "Bukti dummy tanpa data pribadi/rahasia.", "Golden cases scoring yang ditandatangani Process Owner."])
callout(doc, "GATE 6", "Tidak ada defect severity tinggi, semua UAT inti lulus, dan scoring direkonsiliasi sebelum production.", LIGHT_GREEN, GREEN)

# 16
section(doc, 16, "Deployment staging dan production", "Migrasi lebih dulu, deploy terukur, lalu lindungi hostname dengan Access.", False)
heading(doc, "16.1 Deploy staging", 2)
numbers(doc, ["Pastikan konfigurasi staging menunjuk D1/R2 staging.", "Terapkan migrasi remote staging.", "Seed tiga depo dan akun uji tanpa data produksi.", "Build dan deploy Worker + static assets.", "Tambahkan custom domain staging dan Access policy.", "Jalankan smoke test, E2E, dan tes RBAC negatif."])
code(doc, "npx wrangler d1 migrations apply DB --remote\nnpm run build\nnpm run deploy", "STAGING")
heading(doc, "16.2 Checklist pre-production", 2)
bullets(doc, ["☐ Scoring final dan golden case lulus.", "☐ Pengguna production diverifikasi dua pihak.", "☐ Tidak ada data dummy production.", "☐ Bucket private dan CORS terbatas.", "☐ Secret Access production terpasang; tidak di Git.", "☐ Backup dilakukan sebelum migrasi.", "☐ Release notes dan rollback tersedia.", "☐ PIC dan Auditor tahu jadwal cutover."])
heading(doc, "16.3 Deploy production", 2)
code(doc, """npx wrangler whoami
npx wrangler d1 migrations list DB --remote
npx wrangler d1 migrations apply DB --remote
npm run build
npm run deploy""", "PRODUCTION")
numbers(doc, ["Catat commit dan waktu deployment.", "Hubungkan qas.<domain> ke Worker production.", "Buat Access policy production dan salin AUD production.", "Uji satu Admin, satu PIC, satu Auditor, dan satu akses ditolak.", "Jalankan dry run tanpa data sensitif."])
heading(doc, "16.4 Rollback", 2)
bullets(doc, ["Kode: kembali ke deployment/commit sehat.", "Data: gunakan backup/Time Travel D1; rollback kode tidak membalik migrasi data.", "Migrasi: gunakan perubahan kompatibel ke depan.", "R2: jangan hapus massal saat rollback; rekonsiliasi metadata.", "Catat dampak, pengguna, waktu, dan pemulihan."])
callout(doc, "STOP", "Jangan migrasi production bila wrangler whoami, nama database, atau hostname menunjuk staging/akun yang salah.", LIGHT_RED, RED)

# 17
section(doc, 17, "Pilot, pelatihan, dan rollout tiga depo", "Mulai terkendali, ukur masalah, lalu perluas.")
heading(doc, "17.1 Tahap rollout", 2)
table(doc, ["Tahap", "Cakupan", "Kriteria keluar"], [["Pilot", "Karawang; satu siklus contoh", "Self, official, compare, bukti, dashboard, support berjalan"], ["Perluasan 1", "Baros", "Tidak ada defect tinggi; materi pelatihan terbukti"], ["Perluasan 2", "Cirebon", "Hak akses dan data lintas depo benar"], ["Operasional", "Tiga depo", "Owner, backup, monitoring, cadence review aktif"]], [1.2, 2.2, 3.7])
heading(doc, "17.2 Skenario pelatihan PIC", 2)
numbers(doc, ["Login dan cek identitas/depo.", "Buka siklus dan pahami status/tenggat.", "Isi jawaban/catatan dan foto bukti.", "Simulasikan koneksi putus dan lanjutkan draft.", "Review pertanyaan belum lengkap dan submit.", "Setelah official selesai, baca compare dan acknowledge."])
heading(doc, "17.3 Skenario pelatihan Auditor", 2)
numbers(doc, ["Pilih siklus dan depo yang self audit selesai.", "Buktikan jawaban self tidak terlihat.", "Isi audit resmi dan bukti independen.", "Review, submit, dan lihat compare.", "Jelaskan gap/temuan tindak lanjut.", "Gunakan reopen dengan alasan bila koreksi diperlukan."])
heading(doc, "17.4 Form catatan pilot", 2)
table(doc, ["No.", "Temuan", "Dampak", "Owner", "Target", "Status"], [[str(i), "", "", "", "", ""] for i in range(1, 6)], [.45, 2.35, 1.25, 1.1, 1.0, .95])
callout(doc, "GATE 7", "Pemilik proses menyetujui pilot Karawang sebelum akun Baros dan Cirebon diaktifkan untuk siklus produksi.", LIGHT_GREEN, GREEN)

# 18
section(doc, 18, "Operasi, backup, monitoring, dan insiden", "Go-live adalah awal pengelolaan aplikasi.")
heading(doc, "18.1 SOP siklus rutin", 2)
table(doc, ["Kegiatan", "Pelaksana", "Kontrol"], [["Buat/open siklus", "Admin QAS", "Template published dan tanggal benar"], ["Self audit", "PIC depo", "Selesai sebelum due date"], ["Audit resmi", "Auditor", "Setelah self submitted"], ["Review compare", "PIC + Auditor", "Gap utama dijelaskan"], ["Acknowledgement", "PIC", "Waktu dan pengguna tercatat"], ["Close siklus", "Admin QAS", "Tidak ada audit wajib tertinggal"]], [2.0, 1.5, 3.6])
heading(doc, "18.2 Pengelolaan pengguna", 2)
bullets(doc, ["Permintaan mencantumkan email, nama, role, depo, pemohon, persetujuan.", "Gunakan least privilege; Viewer untuk baca saja.", "Saat mutasi/keluar, nonaktifkan user D1 dan keluarkan dari Access bila perlu.", "Tinjau user aktif dan scope secara berkala.", "Jangan hapus user yang menjadi actor audit_events."])
heading(doc, "18.3 Backup dan restore drill", 2)
code(doc, "npx wrangler d1 export qas-audit-prod-db --remote --output backups/qas-audit-prod-YYYYMMDD.sql", "TERMINAL")
bullets(doc, ["Simpan backup terbatas dan catat checksum.", "Ekspor metadata bukti untuk rekonsiliasi R2-D1.", "Uji restore ke non-production.", "Catat RPO/RTO dan hasil latihan."])
heading(doc, "18.4 Monitoring", 2)
table(doc, ["Sinyal", "Yang dipantau", "Respons"], [["Error Worker", "5xx, exception, route gagal", "Periksa request ID dan deployment"], ["D1", "query gagal, latency, growth", "Optimalkan indeks/query atau rollback"], ["R2", "upload/download gagal, orphan", "Periksa binding, ukuran, metadata"], ["Access", "login gagal, policy deny", "Cek email, IdP, hostname, AUD"], ["Bisnis", "audit overdue, gap abnormal", "Hubungi owner depo/Auditor"], ["Kuota", "request, row read/write, storage", "Kurangi query boros/evaluasi plan"]], [1.1, 3.0, 3.0])
heading(doc, "18.5 Prosedur insiden", 2)
numbers(doc, ["Catat waktu, gejala, hostname, pengguna, depo, audit ID, request ID.", "Tentukan dampak: login, data, pengisian, submit, scoring, bukti, dashboard.", "Hentikan perubahan berisiko; jangan hapus bukti/riwayat.", "Pulihkan lewat rollback kode atau perbaikan teruji.", "Rekonsiliasi audit/file selama insiden.", "Dokumentasikan akar masalah, pencegahan, dan komunikasi."])

# 19
section(doc, 19, "Troubleshooting", "Mulai dari identitas, konfigurasi target, lalu data dan UI.", False)
table(doc, ["Gejala", "Pemeriksaan", "Perbaikan"], [["Login berulang", "Hostname, app, policy, cookie, team domain", "Samakan hostname/app; cek IdP"], ["401 Invalid token", "Header, issuer, AUD, waktu", "Secret benar; validasi JWKS/audience"], ["403 setelah login", "Email, active, role, depot", "Perbaiki master/scope"], ["D1 no such table", "Migrasi lokal/remote dan binding", "Apply ke database benar"], ["Data staging muncul", "database_id, bucket, APP_ENV", "Pisahkan config dan redeploy"], ["SPA 404 refresh", "not_found_handling", "single-page-application"], ["Upload gagal", "MIME, size, R2, token, status", "Perbaiki validasi/binding"], ["Foto ganda", "Idempotency dan retry", "Deduplicate operation ID/hash"], ["Draft konflik", "serverVersion vs local", "Resolusi konflik"], ["Skor berbeda", "versi, nilai, bobot, N/A, rounding", "Bandingkan snapshot/golden"], ["Auditor lihat self", "API field filtering", "Tutup endpoint; tes regresi"], ["Dashboard beda", "filter, timezone, void/reopen", "Satu service agregasi"]], [1.45, 2.8, 2.85], 7.5)
heading(doc, "19.1 Urutan diagnosis", 2)
numbers(doc, ["Reproduksi dengan akun/audit uji.", "Salin request ID dan waktu.", "Periksa deployment dan environment.", "Periksa Access identity dan RBAC.", "Periksa API, D1 row, lalu R2 metadata.", "Tambahkan tes regresi sebelum memperbaiki dan deploy."])

# 20
section(doc, 20, "Checklist serah terima dan go-live", "Satu daftar kontrol untuk memastikan produk benar-benar siap digunakan.")
for title, items in [
    ("Bisnis dan data", ["☐ 17 pertanyaan/urutan sesuai sumber.", "☐ Pilihan, bobot, N/A, pembulatan, kategori disahkan.", "☐ Tiga depo tanpa duplikasi.", "☐ Bukti, reopen, acknowledgement, retensi disahkan."]),
    ("Keamanan", ["☐ Access policy sesuai daftar.", "☐ JWT diverifikasi di Worker.", "☐ Tes IDOR/lintas depo lulus.", "☐ R2 private; secret tidak di Git.", "☐ Blind audit terbukti API/UI."]),
    ("Fungsional", ["☐ Self, official, compare, dashboard, ekspor selesai.", "☐ Submit idempotent dan state machine teruji.", "☐ Audit trail/reopen lengkap.", "☐ Skor cocok golden cases."]),
    ("Mobile dan offline", ["☐ Ponsel nyata diuji kamera/offline/sync.", "☐ Draft bertahan setelah refresh.", "☐ Konflik tidak overwrite diam-diam.", "☐ Submit diblokir bila pending."]),
    ("Operasi", ["☐ Panduan per role tersedia.", "☐ Backup/restore drill selesai.", "☐ Monitoring dan insiden aktif.", "☐ Owner dan support ditetapkan.", "☐ Pilot Karawang disetujui."]),
]:
    heading(doc, title, 2); bullets(doc, items)
heading(doc, "20.1 Berita acara go-live", 2)
table(doc, ["Peran", "Nama", "Keputusan / tanda tangan", "Tanggal"], [[r, "", "", ""] for r in ["Process Owner QAS", "Auditor QAS", "PIC QAS Karawang", "PIC QAS Baros", "PIC QAS Cirebon", "Application Owner"]], [1.75, 1.8, 2.65, .9])

# 21
section(doc, 21, "Paket prompt Antigravity", "Salin satu prompt, tinjau hasilnya, lalu lanjut ke prompt berikutnya.")
prompt(doc, 1, "Scaffold dan guardrails", "Baca PRD dan panduan di docs/. Buat fondasi React + TypeScript + Vite PWA dan Cloudflare Worker + Hono. Terapkan struktur folder, routing dasar, error boundary, lint, typecheck, Vitest, dan Playwright. Jangan membangun fitur bisnis atau menebak scoring. Tampilkan rencana file sebelum mengubahnya.", ["dev, build, lint, typecheck, test berhasil.", "Responsif 360 px dan desktop.", "Tidak ada secret/data nyata."])
prompt(doc, 2, "Schema D1 dan seed", "Implementasikan migrasi D1 Bab 8. Tambahkan foreign key, unique/check constraint, dan indeks. Buat seed idempotent tiga depo, akun uji, template QAS draft, J1/J2, dan 17 pertanyaan sumber. Jangan memberi numeric_value, bobot, atau kategori bila belum disahkan. Tambahkan tes schema/seed.", ["Apply lokal dari kosong berhasil.", "Tepat 3 depo dan 17 pertanyaan.", "Tidak ada published tanpa scoring."])
prompt(doc, 3, "Authentication dan RBAC", "Tambahkan middleware Access JWT dari Cf-Access-Jwt-Assertion dengan jose. Validasi issuer, audience, signature, expiry, lalu map email ke users/scopes. Buat /api/me dan helper role/depot. Tes 401, inactive, cross-depot 403, Viewer write 403. Jangan percaya header email tanpa JWT.", ["Tes RBAC positif/negatif lulus.", "Error tidak bocorkan internal.", "Respons memiliki requestId."])
prompt(doc, 4, "Master dan template berversi", "Bangun admin pengguna/scope dan template/version/section/question/option. Published immutable; perubahan via clone draft. Publish memvalidasi scoring lengkap dan tetap diblokir bila belum disahkan. Catat mutasi ke audit_events.", ["CRUD hanya Admin.", "Published tak dapat diubah API/UI.", "Audit trail lengkap."])
prompt(doc, 5, "Siklus dan self audit", "Bangun cycles dan slot SELF/OFFICIAL. Implementasikan self audit responsif, progres, catatan, autosave server, IndexedDB draft, bukti wajib, review, dan submit idempotent dengan role/depot/state/version validation.", ["PIC hanya depo sendiri.", "Draft bertahan refresh/offline.", "Submit tak dapat ganda."])
prompt(doc, 6, "Bukti R2 private", "Implementasikan upload/download lewat Worker ke R2 private. Validasi role, depot, status, MIME, ukuran. Object key acak; metadata/hash D1. Tambahkan retry offline dan tes invalid/large/cross-depot/orphan.", ["Tidak ada URL publik permanen.", "Download cek scope.", "Retry tidak menduplikasi."])
prompt(doc, 7, "Official blind audit", "Bangun official audit setelah self SUBMITTED. Auditor hanya melihat status selesai; payload API/UI tidak boleh berisi jawaban, score, catatan, atau bukti self sampai official SUBMITTED. Tambahkan tes kebocoran semua endpoint.", ["Terkunci sebelum self submit.", "Payload blind bersih.", "Tes regresi lulus."])
prompt(doc, 8, "Scoring dan compare", "Setelah keputusan scoring resmi tersedia, implementasikan engine dari audit_template_versions. Simpan snapshot nilai/config hash. Buat comparison_snapshots. Tambahkan golden tests maksimum, minimum, N/A, ambang, pembulatan, dan contoh manual. Jangan hardcode di UI.", ["Golden cases cocok.", "Snapshot immutable.", "Satu engine untuk API/dashboard/export."])
prompt(doc, 9, "Dashboard dan ekspor", "Bangun dashboard berdasarkan scope dengan filter periode, depo, seksi, status, jenis. Tampilkan status, skor, gap, beda jawaban, temuan prioritas. Ekspor CSV memakai service/filter sama. Mobile berupa card/list.", ["Dashboard/detail/CSV cocok.", "Filter tidak bocor scope.", "360 px tidak rusak."])
prompt(doc, 10, "Hardening dan audit trail", "Tinjau endpoint untuk Zod, RBAC, depot, state, idempotency, concurrency, upload safety, audit_events, rate limit, security headers, dan log terstruktur tanpa data sensitif. Tes IDOR dan tampering.", ["Endpoint write punya tes auth/state.", "Log bersih data sensitif.", "Tak ada defect tinggi."])
prompt(doc, 11, "PWA, offline, dan UX", "Lengkapi manifest, icon, service worker/app shell, IndexedDB, sync queue, conflict handling, status koneksi. Jangan cache audit sensitif secara luas. Uji aksesibilitas, mobile, kamera. Submit nonaktif saat pending.", ["Skenario offline lulus.", "PWA dapat dipasang.", "Draft selamat saat token expire."])
prompt(doc, 12, "Release candidate", "Jalankan lint, typecheck, unit, integration, E2E, build, dan dependency check. Siapkan release notes, migrasi, seed production aman, smoke test, rollback, checklist Bab 16/20. Jangan deploy bila golden scoring, UAT, atau Access/RBAC gagal.", ["Semua gate berbukti.", "Staging smoke lulus.", "Production checklist siap tanda tangan."])

# 22
section(doc, 22, "Referensi resmi", "Gunakan dokumentasi resmi saat versi CLI atau menu dashboard berubah.", False)
refs = [
    ("Cloudflare Workers — Get started", "https://developers.cloudflare.com/workers/get-started/guide/"),
    ("Cloudflare Static Assets", "https://developers.cloudflare.com/workers/static-assets/"),
    ("Cloudflare Vite plugin", "https://developers.cloudflare.com/workers/vite-plugin/get-started/"),
    ("SPA routing", "https://developers.cloudflare.com/workers/static-assets/routing/single-page-application/"),
    ("Wrangler configuration", "https://developers.cloudflare.com/workers/wrangler/configuration/"),
    ("D1 Wrangler commands", "https://developers.cloudflare.com/d1/wrangler-commands/"),
    ("D1 migrations", "https://developers.cloudflare.com/d1/reference/migrations/"),
    ("R2 Workers API", "https://developers.cloudflare.com/r2/get-started/workers-api/"),
    ("Access self-hosted app", "https://developers.cloudflare.com/cloudflare-one/access-controls/applications/http-apps/self-hosted-public-app/"),
    ("Access One-time PIN", "https://developers.cloudflare.com/cloudflare-one/integrations/identity-providers/one-time-pin/"),
    ("Access Validate JWT", "https://developers.cloudflare.com/cloudflare-one/access-controls/applications/http-apps/authorization-cookie/validating-json/"),
    ("Workers pricing", "https://developers.cloudflare.com/workers/platform/pricing/"),
    ("D1 pricing", "https://developers.cloudflare.com/d1/platform/pricing/"),
    ("R2 pricing", "https://developers.cloudflare.com/r2/pricing/"),
    ("Google Antigravity — Getting started", "https://www.antigravity.google/docs/getting-started"),
]
for label, url in refs:
    p = doc.add_paragraph(style="Body Text"); runfmt(p.add_run(label + ": "), 9.1, True, NAVY); hyperlink(p, url, url)
heading(doc, "Catatan sumber internal", 2)
bullets(doc, ["Form Audit QAS.pdf — sumber pertanyaan dan tampilan form eksisting.", "PRD_Aplikasi_Audit_QAS_Cloudflare.docx — kebutuhan produk, arsitektur, data, keamanan, dan acceptance criteria."])
callout(doc, "HASIL AKHIR", "Ikuti Gate 1 sampai Gate 7. Aplikasi siap produksi hanya jika scoring disahkan, UAT lulus, blind audit terbukti, Access/RBAC aman, backup diuji, dan pilot Karawang disetujui.", CREAM, NAVY)

doc.core_properties.title = "Panduan Lengkap Pembangunan Aplikasi Audit QAS"
doc.core_properties.subject = "Cloudflare + Antigravity"
doc.core_properties.author = "Motorcycle Logistic"
doc.core_properties.keywords = "QAS, audit, Cloudflare, Antigravity, D1, R2, PWA"
doc.save(OUT)
print(OUT)

